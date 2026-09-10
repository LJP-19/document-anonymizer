"""Batch processing (spec section 45).

The workflow the UI drives:

    add files or a folder
        -> one tab per document, analysed independently
        -> the user approves, dismisses, or leaves each one
        -> approved documents are processed to the output folder
        -> a summary of what changed, per document

A single shared roster runs through the whole batch, so one client keeps one
pseudonym across every file in it. Processing is cancellable between documents
and reports progress as it goes; a cancelled run leaves the documents it already
wrote in place rather than deleting them.
"""

from __future__ import annotations

import threading
from dataclasses import dataclass, field
from enum import Enum
from pathlib import Path
from typing import Callable, Iterable, Optional

from .decisions.manager import DecisionState
from .entities.roster import ClientRoster
from .session import AnonymizationSession, Status

PDF_GLOB = "*.pdf"


class ItemState(str, Enum):
    PENDING = "Pending"
    ANALYZING = "Analyzing"
    READY = "Ready to review"
    APPROVED = "Approved"
    DISMISSED = "Dismissed"
    PROCESSING = "Processing"
    DONE = "Done"
    FAILED = "Failed"


@dataclass
class Change:
    """One original-to-pseudonym substitution, for the summary table."""

    pii_type: str
    original: str
    pseudonym: str
    occurrences: int
    pages: list[int] = field(default_factory=list)


@dataclass
class BatchItem:
    source_path: str
    session: Optional[AnonymizationSession] = None
    state: ItemState = ItemState.PENDING
    output_path: Optional[str] = None
    mapping_path: Optional[str] = None
    error: str = ""
    verification: Optional[object] = None
    changes: list[Change] = field(default_factory=list)

    @property
    def name(self) -> str:
        return Path(self.source_path).name

    @property
    def detections(self) -> int:
        return len(self.session.candidates) if self.session else 0

    def review_counts(self) -> tuple[int, int]:
        """(reviewed, unreviewed) counted in DISTINCT VALUES, not occurrences.

        The review list groups repeats, so a name appearing 42 times is one row.
        Counting candidates made the prompt say "42 reviewed" beside a chip
        reading 1.
        """
        if not self.session:
            return 0, 0
        reviewed = unreviewed = 0
        for group in self.session.decisions.occurrence_groups(self.session.candidates):
            if all(self.session.decisions.is_reviewed(c) for c in group.candidates):
                reviewed += 1
            else:
                unreviewed += 1
        return reviewed, unreviewed

    @property
    def flagged(self) -> int:
        """Distinct values still awaiting review."""
        return self.review_counts()[1]

    @property
    def _unused_flagged(self) -> int:
        if not self.session:
            return 0
        return len(
            [
                c
                for c in self.session.candidates
                if c.needs_review and not self.session.decisions.is_reviewed(c)
            ]
        )

    @property
    def verified(self) -> bool:
        return bool(self.verification and getattr(self.verification, "passed", False))


@dataclass
class BatchProgress:
    stage: str
    current: int = 0
    total: int = 0
    document: str = ""
    seconds_left: Optional[float] = None

    @property
    def eta(self) -> str:
        """Rough time remaining, from how long the finished documents took."""
        if self.seconds_left is None or self.seconds_left < 1:
            return ""
        seconds = int(self.seconds_left)
        if seconds < 60:
            return f"about {seconds}s left"
        minutes, seconds = divmod(seconds, 60)
        return f"about {minutes}m {seconds:02d}s left"

    @property
    def text(self) -> str:
        parts = [self.stage]
        if self.total:
            parts.append(f"{self.current} of {self.total}")
        if self.eta:
            parts.append(self.eta)
        return "  \u00b7  ".join(parts)


class Cancelled(RuntimeError):
    pass


@dataclass
class Batch:
    """The whole working set."""

    items: list[BatchItem] = field(default_factory=list)
    output_folder: Optional[str] = None
    mapping_folder: Optional[str] = None
    roster_path: Optional[str] = None
    roster: ClientRoster = field(default_factory=ClientRoster)
    changes_workbook: Optional[str] = None
    _cancel: threading.Event = field(default_factory=threading.Event)

    # -- input ------------------------------------------------------------

    def add_files(self, paths: Iterable[str]) -> list[BatchItem]:
        added = []
        existing = {item.source_path for item in self.items}
        for path in paths:
            resolved = str(Path(path).expanduser().resolve())
            if resolved in existing or not resolved.lower().endswith(".pdf"):
                continue
            item = BatchItem(source_path=resolved)
            self.items.append(item)
            added.append(item)
            existing.add(resolved)
        return added

    def add_folder(self, folder: str, recursive: bool = True) -> list[BatchItem]:
        root = Path(folder).expanduser()
        finder = root.rglob if recursive else root.glob
        return self.add_files(sorted(str(p) for p in finder(PDF_GLOB)))

    def remove(self, item: BatchItem) -> None:
        if item in self.items:
            self.items.remove(item)

    def clear(self) -> None:
        """Start over. Files already written to the output folder are untouched."""
        self.items = []
        self._cancel.clear()

    # -- state ------------------------------------------------------------

    @property
    def approved(self) -> list[BatchItem]:
        return [i for i in self.items if i.state is ItemState.APPROVED]

    @property
    def completed(self) -> list[BatchItem]:
        return [i for i in self.items if i.state is ItemState.DONE]

    def approve(self, item: BatchItem) -> None:
        if item.session is not None:
            item.state = ItemState.APPROVED

    def dismiss(self, item: BatchItem) -> None:
        item.state = ItemState.DISMISSED

    def unapprove(self, item: BatchItem) -> None:
        if item.state in (ItemState.APPROVED, ItemState.DISMISSED):
            item.state = ItemState.READY

    # -- analysis ---------------------------------------------------------

    def analyse(
        self, item: BatchItem, progress: Optional[Callable[[BatchProgress], None]] = None
    ) -> BatchItem:
        item.state = ItemState.ANALYZING
        if progress:
            progress(BatchProgress("Reading the document", document=item.name))
        try:
            if self.roster_path:
                self.roster = self.roster or ClientRoster.load(Path(self.roster_path))
            session = AnonymizationSession(
                source_path=item.source_path,
                roster=self.roster,
                roster_path=self.roster_path,
            )
            if progress:
                progress(BatchProgress("Detecting sensitive values", document=item.name))
            session.analyse()
            item.session = session
            item.state = ItemState.READY
        except Exception as exc:  # noqa: BLE001 - surfaced on the item
            item.state = ItemState.FAILED
            item.error = f"{type(exc).__name__}: {exc}"
        return item

    # -- processing -------------------------------------------------------

    def cancel(self) -> None:
        self._cancel.set()

    def reset_cancel(self) -> None:
        self._cancel.clear()

    @property
    def cancelled(self) -> bool:
        return self._cancel.is_set()

    def output_path_for(self, item: BatchItem) -> str:
        session = item.session
        name = session.safe_output_name() if session else f"{Path(item.source_path).stem}.anonymized.pdf"
        folder = Path(self.output_folder) if self.output_folder else Path(item.source_path).parent
        folder.mkdir(parents=True, exist_ok=True)
        candidate = folder / name
        # Never silently overwrite a document already produced.
        counter = 2
        while candidate.exists():
            candidate = folder / f"{Path(name).stem}-{counter}{Path(name).suffix}"
            counter += 1
        return str(candidate)

    def reviewed_counts(self) -> tuple[int, int]:
        """(reviewed, unreviewed) across every approved document."""
        reviewed = unreviewed = 0
        for item in self.approved:
            done, pending = item.review_counts()
            reviewed += done
            unreviewed += pending
        return reviewed, unreviewed

    def keep_unreviewed(self) -> int:
        """Set every unreviewed item to Keep, so only reviewed ones are applied."""
        from .decisions.manager import DecisionState

        changed = 0
        for item in self.approved:
            if not item.session:
                continue
            pending = [
                c for c in item.session.candidates
                if not item.session.decisions.is_reviewed(c)
            ]
            if pending:
                item.session.decisions.set_state(pending, DecisionState.SKIPPED)
                changed += len(pending)
        return changed

    def redact_everything(self) -> int:
        from .decisions.manager import DecisionState

        changed = 0
        for item in self.approved:
            if not item.session:
                continue
            item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
            changed += len(item.session.candidates)
        return changed

    def process_approved(
        self, progress: Optional[Callable[[BatchProgress], None]] = None
    ) -> list[BatchItem]:
        """Export every approved document. Cancellable between documents.

        The cancel flag is NOT cleared here. Clearing it on entry would lose a
        cancel that arrived between the user pressing the button and this call
        starting - a real race in a threaded UI. Callers clear it deliberately
        with `reset_cancel()` when beginning a run.
        """
        import time

        queue = self.approved
        done: list[BatchItem] = []
        started = time.monotonic()
        for index, item in enumerate(queue, start=1):
            elapsed = time.monotonic() - started
            remaining = (
                (elapsed / max(index - 1, 1)) * (len(queue) - index + 1)
                if index > 1 else None
            )
            if self.cancelled:
                break
            item.state = ItemState.PROCESSING
            if progress:
                progress(BatchProgress(
                    f"Pseudonymizing {item.name}", index, len(queue), item.name, remaining
                ))
            try:
                output = self.output_path_for(item)
                if self.mapping_folder:
                    item.session.mapping_folder = self.mapping_folder
                _apply_report, report = item.session.process(output)
                item.output_path = output
                item.mapping_path = item.session.mapping_file
                item.verification = report
                item.changes = summarise_changes(item)
                item.state = ItemState.DONE
                done.append(item)
            except Exception as exc:  # noqa: BLE001
                item.state = ItemState.FAILED
                item.error = f"{type(exc).__name__}: {exc}"
            if progress:
                progress(BatchProgress(
                    f"Verifying {item.name}", index, len(queue), item.name, remaining
                ))
        if done:
            self.changes_workbook = write_changes_workbook(done, self.output_folder)
        if progress:
            progress(BatchProgress("Finished", len(done), len(queue)))
        return done


def write_changes_workbook(items: list[BatchItem], folder: Optional[str]) -> Optional[str]:
    """Save the change table beside the output, one sheet per document.

    It lists originals against their pseudonyms, so it carries client
    identifiers and must be treated exactly like the mapping workbook.
    """
    from pathlib import Path as _Path

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    if not items:
        return None
    base = _Path(folder) if folder else _Path(items[0].source_path).parent
    target = base / "changes" / "what-changed.DO-NOT-SEND.xlsx"
    target.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    workbook.remove(workbook.active)
    for position, item in enumerate(items, start=1):
        title = _Path(item.name).stem[:26] or f"Document {position}"
        sheet = workbook.create_sheet(title=title)
        sheet["A1"] = (
            "CONFIDENTIAL - lists real client values against their pseudonyms. "
            "Never send this with the anonymized document."
        )
        sheet["A1"].font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="C00000")
        sheet["A1"].alignment = Alignment(vertical="center")
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=5)
        sheet.row_dimensions[1].height = 28

        for column, header in enumerate(
            ["Type", "Original", "Pseudonym", "Occurrences", "Pages"], start=1
        ):
            cell = sheet.cell(row=2, column=column, value=header)
            cell.font = Font(name="Arial", size=11, bold=True)
            cell.fill = PatternFill("solid", fgColor="D9D9D9")

        for row, change in enumerate(item.changes, start=3):
            sheet.cell(row=row, column=1, value=change.pii_type.replace("_", " ").title())
            sheet.cell(row=row, column=2, value=change.original)
            sheet.cell(row=row, column=3, value=change.pseudonym)
            sheet.cell(row=row, column=4, value=change.occurrences)
            sheet.cell(row=row, column=5, value=", ".join(str(p) for p in change.pages))
            for column in range(1, 6):
                sheet.cell(row=row, column=column).font = Font(name="Arial", size=11)

        for column, width in enumerate([22, 38, 38, 14, 20], start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = "A3"

    workbook.save(target)
    return str(target)


def summarise_changes(item: BatchItem) -> list[Change]:
    """What actually changed, for the completion view."""
    if not item.session:
        return []
    plan = item.session.plan()
    grouped: dict[tuple[str, str], Change] = {}
    for target in plan.targets:
        key = (target.pii_type.value, target.original.strip())
        change = grouped.get(key)
        if change is None:
            change = Change(
                pii_type=target.pii_type.value,
                original=target.original.strip(),
                pseudonym=target.replacement,
                occurrences=0,
            )
            grouped[key] = change
        change.occurrences += 1
        if target.page_no + 1 not in change.pages:
            change.pages.append(target.page_no + 1)
    return sorted(grouped.values(), key=lambda c: (c.pii_type, c.original))
