"""Batch workflow tests (spec section 45)."""

from __future__ import annotations

from pathlib import Path

import pytest

from app.batch import Batch, ItemState, summarise_changes
from app.decisions.manager import DecisionState
from tests import fixtures


@pytest.fixture
def batch(tmp_path) -> Batch:
    source = tmp_path / "in"
    source.mkdir()
    fixtures.form_pdf(source / "a.pdf")
    fixtures.paragraph_pdf(source / "b.pdf")
    (source / "notes.txt").write_text("not a pdf")
    b = Batch(output_folder=str(tmp_path / "out"))
    b.add_folder(str(source))
    return b


def test_a_folder_adds_only_pdfs(batch):
    assert len(batch.items) == 2
    assert all(i.source_path.endswith(".pdf") for i in batch.items)


def test_adding_the_same_file_twice_is_ignored(batch):
    before = len(batch.items)
    batch.add_files([batch.items[0].source_path])
    assert len(batch.items) == before


def test_approve_dismiss_and_unapprove(batch):
    item = batch.analyse(batch.items[0])
    assert item.state is ItemState.READY
    batch.approve(item)
    assert batch.approved == [item]
    batch.dismiss(item)
    assert not batch.approved
    batch.unapprove(item)
    assert item.state is ItemState.READY


def test_only_approved_documents_are_written(batch, tmp_path):
    first = batch.analyse(batch.items[0])
    batch.analyse(batch.items[1])
    first.session.decisions.set_state(first.session.candidates, DecisionState.ACCEPTED)
    batch.approve(first)

    done = batch.process_approved()
    assert len(done) == 1
    outputs = list((tmp_path / "out").glob("*.pdf"))
    assert len(outputs) == 1


def test_output_folder_is_used_and_never_overwritten(batch, tmp_path):
    item = batch.analyse(batch.items[0])
    item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
    batch.approve(item)
    batch.process_approved()

    batch.unapprove(item)
    batch.approve(item)
    batch.process_approved()

    outputs = sorted(p.name for p in (tmp_path / "out").glob("*.pdf"))
    assert len(outputs) == 2, outputs
    assert outputs[0] != outputs[1]


def test_progress_is_reported(batch):
    item = batch.analyse(batch.items[0])
    item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
    batch.approve(item)
    seen = []
    batch.process_approved(progress=seen.append)
    stages = [p.stage for p in seen]
    assert any(s.startswith("Pseudonymizing") for s in stages), stages
    assert any(s.startswith("Verifying") for s in stages), stages
    assert "Finished" in stages
    # The document being worked on is named, so the user can see what is slow.
    assert any(p.document for p in seen)


def test_cancelling_leaves_finished_documents_in_place(batch, tmp_path):
    for item in batch.items:
        batch.analyse(item)
        item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
        batch.approve(item)
    batch.cancel()
    done = batch.process_approved()
    assert done == []
    assert batch.cancelled


def test_start_over_keeps_written_output(batch, tmp_path):
    item = batch.analyse(batch.items[0])
    item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
    batch.approve(item)
    batch.process_approved()
    written = list((tmp_path / "out").glob("*.pdf"))
    assert written

    batch.clear()
    assert batch.items == []
    assert list((tmp_path / "out").glob("*.pdf")) == written


def test_one_client_keeps_one_pseudonym_across_the_batch(batch):
    for item in batch.items:
        batch.analyse(item)
        item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
        batch.approve(item)
    batch.process_approved()

    names = {}
    for item in batch.completed:
        for change in item.changes:
            if change.original == "John Smith":
                names.setdefault(change.pseudonym, []).append(item.name)
    assert len(names) <= 1, f"the same client got several pseudonyms: {names}"


def test_change_summary_lists_original_and_pseudonym(batch):
    item = batch.analyse(batch.items[0])
    item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
    changes = summarise_changes(item)
    assert changes
    for change in changes:
        assert change.original and change.pseudonym
        assert change.original != change.pseudonym
        assert change.occurrences >= 1
        assert change.pages


def test_reviewed_counts_span_the_approved_documents(batch):
    for item in batch.items:
        batch.analyse(item)
        batch.approve(item)
    reviewed, unreviewed = batch.reviewed_counts()
    assert unreviewed > 0 and reviewed == 0

    first = batch.items[0]
    first.session.decisions.mark_reviewed(first.session.candidates[:2], True)
    reviewed, unreviewed = batch.reviewed_counts()
    assert reviewed == 2


def test_processing_only_reviewed_items_leaves_the_rest_alone(batch, tmp_path):
    """Requested: process what is in Done without deciding the rest."""
    item = batch.analyse(batch.items[0])
    batch.approve(item)
    chosen = item.session.candidates[0]
    item.session.decisions.mark_reviewed([chosen], True)

    kept = batch.keep_unreviewed()
    assert kept == len(item.session.candidates) - 1

    plan = item.session.plan()
    assert [t.candidate_id for t in plan.targets] == [chosen.id]


def test_redact_everything_applies_all_detections(batch):
    item = batch.analyse(batch.items[0])
    batch.approve(item)
    changed = batch.redact_everything()
    assert changed == len(item.session.candidates)
    plan = item.session.plan()
    assert len(plan.targets) == len(item.session.candidates)


def test_counts_are_in_distinct_values_not_occurrences(batch):
    """Reported: '42 reviewed' beside a chip showing 0.

    The review list groups repeats; the counts were tallying every occurrence.
    """
    item = batch.analyse(batch.items[0])
    batch.approve(item)
    groups = item.session.decisions.occurrence_groups(item.session.candidates)
    reviewed, unreviewed = item.review_counts()
    assert reviewed + unreviewed == len(groups)
    assert unreviewed <= len(item.session.candidates)

    item.session.decisions.mark_reviewed(groups[0].candidates, True)
    reviewed, unreviewed = item.review_counts()
    assert reviewed == 1
    assert reviewed + unreviewed == len(groups)


def test_batch_counts_match_the_sum_of_its_documents(batch):
    for item in batch.items:
        batch.analyse(item)
        batch.approve(item)
    expected = [item.review_counts() for item in batch.approved]
    reviewed, unreviewed = batch.reviewed_counts()
    assert reviewed == sum(r for r, _ in expected)
    assert unreviewed == sum(u for _, u in expected)


def test_the_change_table_is_exported_as_excel(batch, tmp_path):
    """Requested: the "what changed" table saved beside the output."""
    from openpyxl import load_workbook

    item = batch.analyse(batch.items[0])
    item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
    batch.approve(item)
    batch.process_approved()

    assert batch.changes_workbook, "no change workbook was written"
    path = Path(batch.changes_workbook)
    assert path.exists()
    assert "DO-NOT-SEND" in path.name, "it lists client values; name it plainly"
    assert path.parent != Path(tmp_path / "out"), "it sits beside the anonymized PDFs"

    sheet = load_workbook(path).worksheets[0]
    assert "CONFIDENTIAL" in str(sheet["A1"].value)
    assert [c.value for c in sheet[2]][:3] == ["Type", "Original", "Pseudonym"]
    rows = [r for r in sheet.iter_rows(min_row=3, values_only=True) if r[0]]
    assert rows and all(r[1] != r[2] for r in rows)


def test_progress_reports_a_time_estimate(batch):
    for item in batch.items:
        batch.analyse(item)
        item.session.decisions.set_state(item.session.candidates, DecisionState.ACCEPTED)
        batch.approve(item)
    seen = []
    batch.process_approved(progress=seen.append)
    # The first document has no history to estimate from; later ones do.
    assert any(p.seconds_left is not None for p in seen), [p.stage for p in seen]

    # An estimate under a second is shown as nothing rather than "about 0s".
    from app.batch import BatchProgress

    assert BatchProgress("x", 1, 2, "d", 45).eta == "about 45s left"
    assert BatchProgress("x", 1, 2, "d", 125).eta == "about 2m 05s left"
    assert BatchProgress("x", 1, 2, "d", 0.2).eta == ""
    assert BatchProgress("x", 1, 2, "d", None).eta == ""
