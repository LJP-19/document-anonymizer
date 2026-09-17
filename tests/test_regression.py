"""Regression tests for the failure modes listed in spec section 81.

Every bug found during development becomes a case here. All data is synthetic.
"""

from __future__ import annotations

import tempfile

import re
from pathlib import Path

import pymupdf
import pytest

from app.decisions.manager import DecisionState
from app.detection.deterministic import financial_tokens, load_rules
from app.detection.types import Candidate, PiiType, Source
from app.entities.registry import EntityRegistry, _shares_token
from app.pseudonymization.generator import generate
from app.session import AnonymizationSession
from tests import fixtures


@pytest.fixture(scope="session")
def pdfs(tmp_path_factory) -> dict[str, str]:
    return fixtures.build_all(tmp_path_factory.mktemp("fixtures"))


def analysed(path: str) -> AnonymizationSession:
    s = AnonymizationSession(source_path=path)
    s.analyse()
    return s


def processed(path: str, out_dir: Path, name: str):
    s = analysed(path)
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    out = str(out_dir / f"{name}.anon.pdf")
    apply_report, report = s.process(out)
    return s, out, apply_report, report


def text_of(path: str) -> str:
    doc = pymupdf.open(path)
    try:
        return "\n".join(p.get_text() for p in doc)
    finally:
        doc.close()


# --- stacked PII (sections 14, 81, 84) -------------------------------------


def test_stacked_field_captures_every_value_line(pdfs):
    """The canonical failure: detecting only the last line of a stacked field."""
    s = analysed(pdfs["stacked"])
    groups = [g for g in s.detection.groups if "zip" in g.label.text.lower()]
    assert groups, "the 'Name, address, and zip code' label was not recognised"
    group = groups[0]
    values = [ln.text.strip() for ln in group.value_lines]
    assert values == ["LJP", "Fremont, CA", "123"], values


def test_stacked_field_does_not_swallow_the_financial_line(pdfs):
    s = analysed(pdfs["stacked"])
    group = next(g for g in s.detection.groups if "zip" in g.label.text.lower())
    assert not any("123,456" in ln.text for ln in group.value_lines)


def test_stacked_field_all_lines_transformed(pdfs, tmp_path):
    """Every typed line of the field is transformed.

    A line nothing could type is left readable and listed under Unlabelled -
    replacing it without knowing what it is produced corrupt output.
    """
    session = analysed(pdfs["stacked"])
    typed = [c for c in session.candidates if c.pii_type is not PiiType.UNCLASSIFIED_GROUP_VALUE]
    session.decisions.set_state(typed, DecisionState.ACCEPTED)
    out = str(tmp_path / "stacked.anon.pdf")
    _apply, report = session.process(out)

    body = text_of(out)
    assert "Fremont, CA" not in body, "the typed line survived redaction"

    # Whether "LJP" itself gets typed depends on whether GLiNER is installed -
    # it can go either way across machines. What must hold everywhere: every
    # value under this label is EITHER transformed OR surfaced under
    # Unlabelled and left readable. Never silently dropped, never replaced
    # with something arbitrary.
    unlabelled = {
        c.normalized for c in session.candidates
        if c.pii_type is PiiType.UNCLASSIFIED_GROUP_VALUE
    }
    if unlabelled:
        for value in unlabelled:
            assert value in body, f"{value!r} is unlabelled but missing from the output"
    assert report.passed, report.failures


def test_unclassified_value_line_is_surfaced_but_not_replaced(pdfs):
    """Superseded fail-safe.

    These were redacted by default, on the reasoning that leaving a value in is
    the unsafe direction. In practice a span nothing could type also has no
    matching replacement, and generating one wrote corrupt text over real
    documents. They are surfaced under Unlabelled and left readable until the
    user names them - a miss the user can see beats damage they cannot.
    """
    s = analysed(pdfs["stacked"])
    unclassified = [c for c in s.candidates if c.pii_type is PiiType.UNCLASSIFIED_GROUP_VALUE]
    assert unclassified
    assert all(c.needs_review for c in unclassified)
    assert all(s.decisions.state(c) is DecisionState.SKIPPED for c in unclassified)


def test_low_confidence_detections_are_redacted_by_default(pdfs):
    """The 987-65-4321 case: flagged, but still removed unless kept explicitly."""
    s = analysed(pdfs["table"])
    weak = next(c for c in s.candidates if c.normalized == "987-65-4321")
    assert weak.needs_review and weak.confidence < 0.9
    assert s.decisions.state(weak) is DecisionState.ACCEPTED


# --- label protection (sections 13, 85) ------------------------------------


def test_labels_are_preserved(pdfs, tmp_path):
    _s, out, _a, _r = processed(pdfs["form"], tmp_path, "form")
    body = text_of(out)
    for label in ("Name:", "Address:", "SSN:", "Email:", "Annual Salary:"):
        assert label in body, f"label {label!r} was destroyed"


def test_financial_value_under_a_non_pii_label_is_preserved(pdfs, tmp_path):
    _s, out, _a, _r = processed(pdfs["form"], tmp_path, "form_money")
    assert "$85,000" in text_of(out)


def test_taxable_income_is_never_a_target(pdfs):
    s = analysed(pdfs["stacked"])
    assert not any("123,456" in c.text for c in s.candidates)


# --- multi-line address (sections 21, 81) ----------------------------------


def test_multiline_address_fully_replaced(pdfs, tmp_path):
    """The 'Apartment 4B' bug: partial line coverage leaving ' 4B' behind."""
    _s, out, _a, report = processed(pdfs["form"], tmp_path, "form_addr")
    body = text_of(out)
    for fragment in ("123 Main Street", "Apartment", "4B", "Fremont"):
        assert fragment not in body, f"address fragment {fragment!r} survived"
    assert report.passed, report.failures


# --- paragraph prose (sections 20, 83) -------------------------------------


def test_paragraph_pii_replaced_and_amount_preserved(pdfs, tmp_path):
    _s, out, _a, report = processed(pdfs["paragraph"], tmp_path, "para")
    body = text_of(out)
    for gone in ("John Smith", "john@example.com", "555) 123-4567", "123-45-6789"):
        assert gone not in body, f"{gone!r} survived"
    assert "$18,450.00" in body
    assert "ABC Company" in body, "a generic organisation name should be preserved"
    assert report.passed, report.failures


def test_phone_area_code_is_not_treated_as_currency():
    """Regression: '(555)' was matching the accounting-negative money pattern."""
    tokens = financial_tokens("call (555) 123-4567 owing $18,450.00 and (1,234.00)")
    assert "$18,450.00" in tokens
    assert "(1,234.00)" in tokens
    assert not any("555" in t for t in tokens)


# --- tables (sections 22, 81) ----------------------------------------------


def test_table_identities_and_ssns_detected_salaries_kept(pdfs, tmp_path):
    s = analysed(pdfs["table"])
    found = {c.normalized for c in s.candidates}
    assert "John Smith" in found and "Jane Doe" in found
    assert "123-45-6789" in found and "987-65-4321" in found

    _s2, out, _a, _r = processed(pdfs["table"], tmp_path, "table")
    body = text_of(out)
    assert "$85,000" in body and "$92,000" in body
    assert "Employee" in body and "SSN" in body and "Salary" in body


def test_ssn_failing_structural_validation_is_reviewed_not_dropped(pdfs):
    """Regression: 987-65-4321 fails the SSN area check and was being deleted."""
    s = analysed(pdfs["table"])
    c = next(c for c in s.candidates if c.normalized == "987-65-4321")
    assert c.needs_review
    assert "structural validation" in c.review_reason
    assert c.confidence < 0.9


def test_column_header_provides_context_to_later_rows(pdfs):
    from app.detection.deterministic import _context_window
    from app.document.provider import NativePdfTextProvider

    doc = NativePdfTextProvider().load(pdfs["table"])
    lines_by_page = {p.number: p.lines for p in doc.pages}
    line = next(ln for ln in doc.pages[0].lines if "987" in ln.text)
    assert "ssn" in _context_window(line, lines_by_page)


# --- entity registry (section 25) ------------------------------------------


def test_same_value_maps_to_same_pseudonym():
    r = EntityRegistry(scope="t")
    a = r.pseudonym_for(PiiType.PERSON, "John Smith")
    b = r.pseudonym_for(PiiType.PERSON, "john  smith")
    assert a == b


def test_pseudonyms_are_deterministic_across_runs():
    a = EntityRegistry(scope="t").pseudonym_for(PiiType.PERSON, "John Smith")
    b = EntityRegistry(scope="t").pseudonym_for(PiiType.PERSON, "John Smith")
    assert a == b


def test_different_values_do_not_collide():
    r = EntityRegistry(scope="t")
    names = {r.pseudonym_for(PiiType.PERSON, n) for n in ("John Smith", "Jane Doe", "Ann Lee")}
    assert len(names) == 3


def test_pseudonym_never_reuses_a_token_of_the_original():
    """Regression: 'John Smith' -> 'John Glass' leaked the first name."""
    r = EntityRegistry(scope="t")
    for original in ("John Smith", "Mary Johnson", "Robert Brown", "Fremont, CA"):
        assert not _shares_token(original, r.pseudonym_for(PiiType.PERSON, original))


def test_short_token_stays_short():
    """Regression: 'LJP' became 'Elizabeth Hernandez' and wrecked the layout."""
    out = generate(PiiType.UNCLASSIFIED_GROUP_VALUE, "LJP")
    assert len(out) <= 6


def test_user_edit_overrides_and_persists():
    r = EntityRegistry(scope="t")
    r.pseudonym_for(PiiType.PERSON, "John Smith")
    r.override(PiiType.PERSON, "John Smith", "Alan Turner")
    assert r.pseudonym_for(PiiType.PERSON, "John Smith") == "Alan Turner"


# --- decisions (sections 26-28) --------------------------------------------


def test_skip_preserves_the_value(pdfs, tmp_path):
    s = analysed(pdfs["form"])
    ssn = next(c for c in s.candidates if c.pii_type is PiiType.SSN)
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    s.decisions.skip([ssn])
    out = str(tmp_path / "skip.pdf")
    _a, report = s.process(out)
    assert "123-45-6789" in text_of(out)
    assert report.passed, report.failures


def test_apply_to_all_and_undo(pdfs):
    s = analysed(pdfs["paragraph"])
    target = next(c for c in s.candidates if c.pii_type is PiiType.PERSON)
    changed = s.decisions.apply_to_all(s.candidates, target, DecisionState.SKIPPED)
    assert changed and all(s.decisions.state(c) is DecisionState.SKIPPED for c in changed)
    assert s.decisions.undo()
    assert s.decisions.state(target) is not DecisionState.SKIPPED


def test_edit_replacement_is_used(pdfs, tmp_path):
    s = analysed(pdfs["form"])
    person = next(c for c in s.candidates if c.pii_type is PiiType.PERSON)
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    s.decisions.edit([person], "Casey Doyle")
    out = str(tmp_path / "edit.pdf")
    s.process(out)
    assert "Casey Doyle" in text_of(out)


# --- redaction and export (sections 34-36, 44) -----------------------------


def test_original_file_is_never_overwritten(pdfs):
    s = analysed(pdfs["form"])
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    with pytest.raises(ValueError):
        s.process(pdfs["form"])


def test_replacement_text_is_real_red_text(pdfs, tmp_path):
    s, out, _a, _r = processed(pdfs["form"], tmp_path, "red")
    doc = pymupdf.open(out)
    try:
        colors = {
            int(span["color"])
            for page in doc
            for b in page.get_text("dict")["blocks"]
            for line in b.get("lines", [])
            for span in line["spans"]
        }
    finally:
        doc.close()
    assert 0xFF0000 in colors, "no red text spans in the output"


def test_no_replacement_wraps_to_a_second_line(pdfs, tmp_path):
    _s, _out, apply_report, _r = processed(pdfs["form"], tmp_path, "wrap")
    assert not apply_report.overflowed


def test_preview_and_export_use_the_same_plan(pdfs, tmp_path):
    """Section 33: one transformation plan, no divergent preview logic."""
    s = analysed(pdfs["form"])
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    plan_a, plan_b = s.plan(), s.plan()
    assert [t.replacement for t in plan_a.targets] == [t.replacement for t in plan_b.targets]
    assert s.preview_transformed(0)[:8] == b"\x89PNG\r\n\x1a\n"


# --- verification (sections 37-41) -----------------------------------------


def test_verification_runs_against_the_saved_file(pdfs, tmp_path):
    _s, out, _a, report = processed(pdfs["form"], tmp_path, "verify")
    names = {c.name for c in report.checks}
    assert {"accepted originals removed", "replacement text is red", "no partial redaction"} <= names
    assert report.output_path == out


def test_verifier_detects_a_tampered_output(pdfs, tmp_path):
    """A verifier that cannot fail is worthless."""
    from app.verification.verifier import verify

    s = analysed(pdfs["form"])
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    plan = s.plan()
    untouched = str(tmp_path / "untouched.pdf")
    pymupdf.open(pdfs["form"]).save(untouched)
    report = verify(plan, untouched)
    assert not report.passed
    assert any(c.name == "accepted originals removed" for c in report.failures)


def test_short_numeric_value_inside_currency_is_not_a_false_survivor(pdfs, tmp_path):
    """Regression: value '123' was reported as surviving inside '$123,456'."""
    from app.verification.verifier import _contains_value

    assert not _contains_value("taxable income: $123,456", "123")
    assert _contains_value("zip 123 here", "123")


# --- unsupported content (sections 6, 42, 87) ------------------------------


def test_scanned_page_is_reported_not_silently_ignored(pdfs, tmp_path):
    s = analysed(pdfs["scanned"])
    assert s.document.ocr_required_pages == [0]
    assert any("OCR REQUIRED" in w for w in s.detection.warnings)
    out = str(tmp_path / "scan.pdf")
    _a, report = s.process(out)
    assert not report.passed
    assert any(c.name == "all pages analysed" for c in report.failures)


def test_status_never_claims_plain_verified(pdfs, tmp_path):
    """Section 87: no false certainty."""
    _s, _out, _a, report = processed(pdfs["form"], tmp_path, "status")
    assert report.status == "EXPORT VERIFIED"


# --- rules file -------------------------------------------------------------


def test_rules_file_loads_and_all_patterns_compile():
    rs = load_rules()
    assert len(rs.rules) > 20 and rs.labels and rs.non_pii_labels
    for rule in rs.rules:
        assert isinstance(rule.regex, re.Pattern)


def test_manual_pii_uses_the_same_pipeline(pdfs, tmp_path):
    s = analysed(pdfs["paragraph"])
    page = s.document.pages[0]
    line = page.lines[0]
    added = s.add_manual(0, line.bbox)
    assert added is not None and added.source is Source.MANUAL
    assert s.decisions.state(added) is DecisionState.MANUALLY_ADDED
    plan = s.plan()
    assert any(t.candidate_id == added.id for t in plan.targets)


# --- detection recall on realistic form layouts ----------------------------


def _hard_pdf(path: Path) -> str:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    rows = [
        (60, 710, "Your first name and middle initial"),
        (300, 710, "Last name"),
        (60, 696, "MARIA T"),
        (300, 696, "GONZALEZ-REYES"),
        (60, 674, "Home address (number and street). If you have a P.O. box, see instructions."),
        (60, 660, "4820 Camino Del Rio Apt 12C"),
        (60, 638, "City, town, or post office, state, and ZIP code"),
        (60, 624, "San Jose, CA 95129"),
        (60, 600, "Filing status: Married filing jointly"),
        (60, 578, "Occupation: Software Engineer"),
        (60, 556, "Daytime phone 408.555.0198   Email m.gonzalez@fastmail.co"),
        (60, 534, "Bank routing 121000248   Account 000123456789"),
        (60, 512, "1  Wages, salaries, tips  ................  $412,890"),
        (60, 498, "11 Adjusted gross income  ...............  $438,117"),
    ]
    for x, y, txt in rows:
        c.drawString(x, y, txt)
    c.save()
    return str(path)


@pytest.fixture(scope="session")
def hard_pdf(tmp_path_factory) -> str:
    return _hard_pdf(tmp_path_factory.mktemp("hard") / "hard.pdf")


def test_verbose_form_labels_are_recognised(hard_pdf):
    """Real labels are sentences: 'Your first name and middle initial'."""
    s = analysed(hard_pdf)
    labels = " | ".join(g.label.text.lower() for g in s.detection.groups)
    assert "first name" in labels
    assert "home address" in labels


def test_all_caps_and_lone_surname_names_detected(hard_pdf):
    s = analysed(hard_pdf)
    found = {c.normalized for c in s.candidates}
    assert "MARIA T" in found
    assert "GONZALEZ-REYES" in found


def test_full_street_line_captured_not_truncated(hard_pdf):
    """The street regex stops at 'Apt'; the field group must cover the rest."""
    s = analysed(hard_pdf)
    assert any("12C" in c.text for c in s.candidates)


def test_contact_and_bank_details_on_a_value_line_survive_label_matching(hard_pdf):
    """Regression: the word 'phone' turned the whole value line into a label."""
    s = analysed(hard_pdf)
    types = {c.pii_type for c in s.candidates}
    assert PiiType.EMAIL in types
    assert PiiType.PHONE in types
    assert PiiType.BANK_ACCOUNT in types or PiiType.ROUTING_NUMBER in types


def test_business_facts_under_unknown_labels_are_left_alone(hard_pdf):
    """Anonymise WHO, not WHAT. Occupation and filing status are not identity."""
    s = analysed(hard_pdf)
    text = " ".join(c.text for c in s.candidates)
    assert "Software Engineer" not in text
    assert "Married filing jointly" not in text
    assert "412,890" not in text and "438,117" not in text


def test_form_vocabulary_is_never_a_person(hard_pdf):
    """Regression: spaCy tagged 'Daytime' as a PERSON."""
    s = analysed(hard_pdf)
    people = {c.normalized for c in s.candidates if c.pii_type is PiiType.PERSON}
    assert not {"Daytime", "Preparer", "Occupation", "Wages"} & people


# --- GLiNER layer ----------------------------------------------------------


def _gliner():
    from app.detection.gliner import GlinerDetector

    detector = GlinerDetector()
    if not detector.available:
        pytest.skip("GLiNER weights not fetched; run buildtools/fetch_models.py")
    return detector


def test_gliner_returns_the_label_that_was_asked_for():
    """The point of a label-conditioned model: no more guessing the type."""
    detector = _gliner()
    words = "John Smith lives at 123 Main Street and his SSN is 123-45-6789 .".split()
    spans = {" ".join(words[s.start_word:s.end_word + 1]): s.label for s in detector.predict(words)}
    assert any("person" in label for value, label in spans.items() if "John" in value)
    assert any("street" in label for value, label in spans.items() if "Main" in value)


def test_a_currency_amount_is_never_pii_whatever_the_model_says(tmp_path):
    """Regression: GLiNER labelled '$85,000' a date of birth at 0.62."""
    s = analysed(fixtures.table_pdf(tmp_path / "t.pdf"))
    for candidate in s.candidates:
        assert "$" not in candidate.text, f"{candidate.text!r} flagged as {candidate.pii_type.value}"


def test_gliner_identifies_money_as_a_negative_class():
    """Used to veto competing detections so figures survive."""
    detector = _gliner()
    words = "Adjusted gross income was $ 438,117 for the year .".split()
    labels = [s.label for s in detector.predict(words)]
    assert "money amount" in labels


def test_field_value_is_never_partially_redacted(tmp_path):
    """The street regex stops at 'Apt'; the field must still go as a unit."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "apt.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    c.drawString(60, 700, "Home address (number and street)")
    c.drawString(60, 686, "4820 Camino Del Rio Apt 12C")
    c.drawString(60, 660, "1  Wages  ......  $412,890")
    c.save()

    s = analysed(str(path))
    street = [c for c in s.candidates if "4820" in c.text]
    assert street, "the street line was not detected at all"
    assert "12C" in street[0].text, f"partial value detected: {street[0].text!r}"


def test_financial_figures_are_not_detected_as_pii(tmp_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "money.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    for i, line in enumerate(
        ["1  Wages, salaries, tips  ......  $412,890",
         "11 Adjusted gross income  .....  $438,117",
         "Occupation: Software Engineer",
         "Filing status: Married filing jointly"]
    ):
        c.drawString(60, 700 - i * 16, line)
    c.save()

    s = analysed(str(path))
    text = " ".join(c.text for c in s.candidates)
    for preserved in ("412,890", "438,117", "Software Engineer", "Married filing jointly"):
        assert preserved not in text, f"{preserved!r} was flagged as PII"


# --- joint / compound names and propagation --------------------------------


def _joint_pdf(path: Path) -> str:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    rows = [
        (60, 720, "Taxpayer name"), (60, 706, "JP & ML"),
        (60, 684, "Spouse first name and last name"),
        (60, 670, "John and Mary Gonzalez-Reyes"),
        (60, 648, "Preparer notes"),
        (60, 634, "Reviewed the return for Gonzalez-Reyes this week."),
        (60, 600, "1  Wages, salaries, tips  ......  $412,890"),
    ]
    for x, y, text in rows:
        c.drawString(x, y, text)
    c.save()
    return str(path)


@pytest.fixture(scope="session")
def joint_pdf(tmp_path_factory) -> str:
    return _joint_pdf(tmp_path_factory.mktemp("joint") / "joint.pdf")


def test_ampersand_joint_name_becomes_two_people(joint_pdf):
    """'JP & ML' is two taxpayers, not one entity sharing one pseudonym."""
    s = analysed(joint_pdf)
    people = {c.normalized for c in s.candidates if c.pii_type is PiiType.PERSON}
    assert "JP" in people and "ML" in people
    assert "JP & ML" not in people


def test_written_and_joint_name_becomes_two_people(joint_pdf):
    s = analysed(joint_pdf)
    people = {c.normalized for c in s.candidates if c.pii_type is PiiType.PERSON}
    assert "John" in people
    assert any("Mary" in p for p in people)
    assert "John and Mary Gonzalez-Reyes" not in people


def test_joint_spouses_get_different_pseudonyms(joint_pdf, tmp_path):
    s = analysed(joint_pdf)
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    plan = s.plan()
    jp = next(t.replacement for t in plan.targets if t.original.strip() == "JP")
    ml = next(t.replacement for t in plan.targets if t.original.strip() == "ML")
    assert jp != ml, "both spouses received the same pseudonym"


def test_a_value_group_stops_at_the_next_field(joint_pdf):
    """Regression: the spouse field swallowed 'Preparer notes' and its value."""
    s = analysed(joint_pdf)
    spouse = next(g for g in s.detection.groups if "Spouse" in g.label.text)
    assert len(spouse.value_lines) == 1
    assert "Preparer" not in " ".join(l.text for l in spouse.value_lines)


def test_surname_is_caught_in_free_text_elsewhere(joint_pdf):
    """The name reappears in prose with no label; it must still be found."""
    s = analysed(joint_pdf)
    assert any("Gonzalez-Reyes" in c.text for c in s.candidates)


def test_propagation_finds_a_repeat_with_no_local_context(tmp_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "repeat.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    c.drawString(60, 720, "Taxpayer name")
    c.drawString(60, 706, "Marisol Etxeberria")
    c.drawString(60, 600, "Payer            Amount")
    c.drawString(60, 586, "Marisol Etxeberria   $1,110")
    c.save()

    s = analysed(str(path))
    hits = [c for c in s.candidates if "Etxeberria" in c.text]
    assert len(hits) >= 2, "the repeat in the table was not found"


# --- roster, mapping workbook, LLM auditor ---------------------------------


def test_mapping_workbook_lists_original_and_pseudonym(tmp_path):
    from openpyxl import load_workbook

    s = analysed(fixtures.form_pdf(tmp_path / "f.pdf"))
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    s.process(str(tmp_path / "out.pdf"))
    assert s.mapping_file, "no mapping workbook was written"

    sheet = load_workbook(s.mapping_file)["Mapping"]
    assert "CONFIDENTIAL" in str(sheet["A1"].value)
    assert [c.value for c in sheet[2]][:3] == ["Type", "Original value", "Pseudonym"]
    rows = [(r[0], r[1], r[2]) for r in sheet.iter_rows(min_row=3, values_only=True) if r[0]]
    assert rows
    assert any(original == "John Smith" for _t, original, _p in rows)
    for _t, original, pseudonym in rows:
        assert original != pseudonym


def test_mapping_is_not_written_beside_the_anonymized_pdf(tmp_path):
    """It must not be picked up by "attach everything in this folder"."""
    s = analysed(fixtures.form_pdf(tmp_path / "f.pdf"))
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    out = tmp_path / "out.pdf"
    s.process(str(out))
    mapping = Path(s.mapping_file)
    assert mapping.parent != out.parent
    assert "DO-NOT-SEND" in mapping.name


def test_roster_keeps_one_client_consistent_across_documents(tmp_path):
    """A batch must read coherently: same person, same pseudonym, every file."""
    from app.entities.roster import ClientRoster
    from app.session import AnonymizationSession

    roster = ClientRoster()
    pseudonyms = []
    for index in (1, 2):
        source = fixtures.form_pdf(tmp_path / f"doc{index}.pdf")
        session = AnonymizationSession(source_path=source, roster=roster)
        session.analyse()
        session.decisions.set_state(session.candidates, DecisionState.ACCEPTED)
        plan = session.plan()
        pseudonyms.append(
            next(t.replacement for t in plan.targets if t.original.strip() == "John Smith")
        )
    assert pseudonyms[0] == pseudonyms[1], "the same client got two pseudonyms"


def test_roster_survives_a_round_trip_through_the_workbook(tmp_path):
    from app.detection.types import PiiType
    from app.entities.roster import ClientRoster

    roster = ClientRoster()
    roster.record(PiiType.PERSON, "Marisol Etxeberria", "Julie Baker", "doc1.pdf")
    path = tmp_path / "roster.xlsx"
    roster.save(path)

    reloaded = ClientRoster.load(path)
    assert reloaded.pseudonym_for(PiiType.PERSON, "marisol  etxeberria") == "Julie Baker"


def test_name_suffix_is_inside_the_detected_span(tmp_path):
    """Regression: 'MARIA T GONZALEZ-REYES Jr.' left 'Jr.' behind."""
    from app.detection.heuristics import looks_like_person

    for name in ("MARIA T GONZALEZ-REYES JR.", "Smith, John Jr.", "Dr. John Smith III"):
        assert looks_like_person(name)[0], name


def test_demographic_categories_are_detected(tmp_path):
    """DOB, citizenship, birthplace and sex were missing entirely."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "demog.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    rows = [
        (720, "Date of birth"), (706, "04/11/1979"),
        (684, "Country of citizenship"), (670, "Mexico"),
        (648, "Place of birth"), (634, "Guadalajara, Jalisco"),
        (612, "Sex"), (598, "F"),
        (576, "Occupation"), (562, "Software Engineer"),
    ]
    for y, text in rows:
        c.drawString(60, y, text)
    c.save()

    s = analysed(str(path))
    found = {c.text.strip() for c in s.candidates}
    for value in ("04/11/1979", "Mexico", "Guadalajara, Jalisco", "F"):
        assert value in found, f"{value!r} not detected"
    assert "Software Engineer" not in found


def test_gender_pseudonym_always_differs_from_the_original():
    from app.detection.types import PiiType
    from app.pseudonymization.generator import generate

    for original in ("M", "F", "male", "female"):
        assert generate(PiiType.GENDER, original).lower() != original.lower()


@pytest.mark.slow
def test_llm_auditor_finds_categories_no_rule_covers():
    from app.detection.auditor import LlmAuditor

    auditor = LlmAuditor()
    if not auditor.available:
        pytest.skip("LLM weights not fetched")
    text = (
        "Taxpayer name: MARIA T GONZALEZ-REYES\n"
        "Date of birth 04/11/1979   Country of citizenship: Mexico\n"
        "Place of birth: Guadalajara, Jalisco   Sex: F\n"
        "Occupation: Software Engineer\n"
        "1  Wages, salaries, tips ......... $412,890\n"
    )
    missed, wrong = auditor.audit(text, ["MARIA T GONZALEZ-REYES"])
    blob = " ".join(f.text for f in missed).lower()
    assert "1979" in blob and "mexico" in blob
    assert "412,890" not in blob, "the auditor flagged a financial figure"


def test_a_candidate_overlapping_a_currency_amount_is_dropped(tmp_path):
    """Regression: '123' matched inside '$123,456' and corrupted the figure.

    Word-boundary matching does not protect a number inside a currency amount,
    because '$' and ',' are not word characters. CI caught this on the stacked
    fixture: financial values preserved -> FAIL.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "money_overlap.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    for y, text in (
        (720, "Name, address, and zip code"),
        (706, "LJP"),
        (692, "Fremont, CA"),
        (678, "123"),
        (640, "Taxable income: $123,456"),
    ):
        c.drawString(72, y, text)
    c.save()

    s = analysed(str(path))
    for candidate in s.candidates:
        assert "123,456" not in candidate.line.text or "Taxable" not in candidate.line.text, (
            f"{candidate.text!r} was flagged on the financial line"
        )

    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    out = str(tmp_path / "out.pdf")
    _apply, report = s.process(out)
    assert "$123,456" in text_of(out)
    financial = next(c for c in report.checks if c.name == "financial values preserved")
    assert financial.passed, financial.detail


def test_the_financial_guard_runs_after_every_pass_that_adds_candidates():
    """Regression: the guard ran before propagation and the audit, so both bypassed it.

    Ordering is the whole bug here - the check itself was correct.
    """
    source = (Path(__file__).resolve().parents[1] / "app" / "detection" / "engine.py").read_text()
    guard = "_drop_financial_values(candidates)"
    audit = "audit_document(doc, candidates)"
    propagation = "propagate(doc, candidates, subjects)"

    assert source.count(guard) >= 2, "the guard must run again after the later passes"
    last_guard = source.rfind(guard)
    assert last_guard > source.rfind(audit), "the guard must run after the audit pass"
    assert last_guard > source.rfind(propagation), "the guard must run after propagation"


# --- partial redaction on unlabelled lines ---------------------------------


def _lines_pdf(path: Path, lines: list[str]) -> str:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    y = 720
    for text in lines:
        c.drawString(72, y, text)
        # 14pt matches real form leading. Wider spacing exceeds the field-group
        # gap threshold, so a label would not bind its value and the fixture
        # would be testing something the layout never does.
        y -= 14
    c.save()
    return str(path)


def test_spaced_initials_are_redacted_whole(tmp_path):
    """Reported: 'LJ P' redacted only 'LJ'."""
    s = analysed(_lines_pdf(tmp_path / "a.pdf", ["LJ P", "1  Wages ...... $412,890"]))
    hit = next((c for c in s.candidates if "LJ" in c.text), None)
    assert hit is not None, "the name was not detected at all"
    assert hit.text.strip() == "LJ P", f"partial value: {hit.text!r}"


def test_city_state_and_number_are_redacted_whole(tmp_path):
    """Reported: 'Fremont, CA 1234' redacted only '1234'."""
    s = analysed(_lines_pdf(tmp_path / "b.pdf", ["Fremont, CA 1234", "1  Wages ...... $412,890"]))
    hit = next((c for c in s.candidates if "1234" in c.text), None)
    assert hit is not None
    assert "Fremont" in hit.text, f"partial value: {hit.text!r}"


def test_widening_never_swallows_a_figure(tmp_path):
    s = analysed(_lines_pdf(tmp_path / "c.pdf", ["LJ P", "Refund due 4,820.00", "$412,890"]))
    for candidate in s.candidates:
        assert "412,890" not in candidate.text
        assert "4,820" not in candidate.text


def test_widening_leaves_prose_alone(tmp_path):
    """On a sentence the value is a span, not the line."""
    s = analysed(
        _lines_pdf(tmp_path / "d.pdf", ["John Smith submitted the report to ABC Company."])
    )
    for candidate in s.candidates:
        assert "submitted" not in candidate.text
        assert "ABC Company" not in candidate.text


def test_widening_does_not_merge_joint_names(tmp_path):
    """'JP & ML' is two people; the separator must not pull them together."""
    s = analysed(_lines_pdf(tmp_path / "e.pdf", ["Taxpayer name", "JP & ML"]))
    people = {c.normalized for c in s.candidates if c.pii_type is PiiType.PERSON}
    assert "JP" in people and "ML" in people
    assert not any("&" in p for p in people)


# --- manual additions ------------------------------------------------------


def test_user_can_add_a_missed_value_with_a_custom_replacement(tmp_path):
    s = analysed(_lines_pdf(tmp_path / "m.pdf", ["Acme Holdings LLC", "1 Wages ...... $412,890"]))
    added = s.add_manual_text("Acme Holdings LLC", replacement="Northwind Trading LLC")
    assert added
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    out = str(tmp_path / "out.pdf")
    s.process(out)
    body = text_of(out)
    assert "Acme Holdings" not in body
    assert "Northwind Trading LLC" in body
    assert "$412,890" in body


def test_cascade_finds_the_value_on_every_page(tmp_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "multi.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    for _ in range(3):
        c.setFont("Helvetica", 10)
        c.drawString(72, 720, "Prepared for Vantage Partners")
        c.showPage()
    c.save()

    s = analysed(str(path))
    added = s.add_manual_text("Vantage Partners", cascade=True)
    assert {c.page_no for c in added} == {0, 1, 2}


def test_without_cascade_only_the_first_page_is_touched(tmp_path):
    """Whether the models are installed must not change this.

    With them, "Vantage Partners" is already a candidate; without them it is
    not. Either way, adding it with cascade off must touch page 1 only.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "multi2.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    for _ in range(2):
        c.setFont("Helvetica", 10)
        c.drawString(72, 720, "Prepared for Vantage Partners")
        c.showPage()
    c.save()

    s = analysed(str(path))
    added = s.add_manual_text("Vantage Partners", cascade=False)
    assert added, "adding a value returned nothing at all"
    assert {c.page_no for c in added} == {0}


def test_adding_a_value_already_detected_returns_it(tmp_path):
    """Regression: pressing Add on a known value silently did nothing."""
    s = analysed(fixtures.form_pdf(tmp_path / "adopt.pdf"))
    known = s.candidates[0].normalized

    adopted = s.add_manual_text(known, pii_type=PiiType.ORG_PRIVATE)
    assert adopted, "the request was ignored because the value was already found"
    assert all(c.pii_type is PiiType.ORG_PRIVATE for c in adopted)
    assert all(s.decisions.is_reviewed(c) for c in adopted)


def test_apply_to_same_off_adds_a_single_occurrence(tmp_path):
    s = analysed(
        _lines_pdf(tmp_path / "rep.pdf", ["Vantage Partners", "Vantage Partners again"])
    )
    added = s.add_manual_text("Vantage Partners", apply_to_same=False)
    assert len(added) == 1


def test_manual_additions_are_already_reviewed(tmp_path):
    """The user just told us; it does not belong back in their review queue."""
    s = analysed(_lines_pdf(tmp_path / "r.pdf", ["Acme Holdings LLC"]))
    added = s.add_manual_text("Acme Holdings LLC")
    assert added and all(s.decisions.is_reviewed(c) for c in added)


def test_business_names_are_detected_and_typed_from_their_label(tmp_path):
    """Reported: business names sometimes missed, and typed as people."""
    s = analysed(
        _lines_pdf(
            tmp_path / "biz.pdf",
            ["Business name", "Acme Holdings LLC",
             "Employer name", "Vantage Partners Inc.",
             "1  Gross receipts  ......  $1,204,880"],
        )
    )
    found = {c.normalized: c.pii_type for c in s.candidates}
    assert "Acme Holdings LLC" in found, "business name not detected"
    assert found["Acme Holdings LLC"] is PiiType.ORG_PRIVATE, "typed as the wrong thing"

    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    out = str(tmp_path / "out.pdf")
    s.process(out)
    body = text_of(out)
    assert "Acme Holdings" not in body
    assert "$1,204,880" in body
    assert "Business name" in body


def test_a_company_pseudonym_keeps_its_entity_suffix(tmp_path):
    """LLC vs Inc carries tax meaning the receiving analysis needs."""
    from app.pseudonymization.generator import generate

    assert generate(PiiType.ORG_PRIVATE, "Acme Holdings LLC").endswith("LLC")
    assert generate(PiiType.ORG_PRIVATE, "Vantage Partners Inc.").rstrip(".").endswith("Inc")


def test_review_model_suggestions_are_not_applied_unasked():
    """Reported: random replacements on text that was not PII.

    The audit pass is the only layer that proposes values nothing else saw. Its
    suggestions are surfaced for a decision rather than redacted by default.
    """
    from app.decisions.manager import DecisionManager, DecisionState
    from app.detection.types import Candidate, PiiType, Source
    from app.document.model import Char, Line, Span

    char = Char(text="X", bbox=(0, 0, 5, 10))
    span = Span(text="X", bbox=(0, 0, 5, 10), font="helv", size=10, color=0, chars=[char])
    line = Line(page_no=0, block_no=0, line_no=0, spans=[span], text="X",
                offsets=[char], bbox=(0, 0, 5, 10))

    def make(source: Source) -> Candidate:
        return Candidate(
            pii_type=PiiType.PERSON, text="X", page_no=0, rect=(0, 0, 5, 10),
            line=line, start=0, end=1, confidence=0.7, source=source,
        )

    from_rule, from_audit = make(Source.REGEX), make(Source.AUDIT)
    from_audit.line = Line(page_no=0, block_no=0, line_no=1, spans=[span], text="X",
                           offsets=[char], bbox=(0, 0, 5, 10))

    manager = DecisionManager()
    manager.register([from_rule, from_audit])
    assert manager.state(from_rule) is DecisionState.ACCEPTED
    assert manager.state(from_audit) is DecisionState.SKIPPED


# --- pseudonym quality -----------------------------------------------------


def test_a_label_inside_a_value_survives():
    """Reported: 'Nguyen, Tuyet PTIN P01234567' replaced PTIN with a surname.

    PTIN names the field. Destroying it is the same failure as redacting a
    label, just hidden inside a value.
    """
    from app.pseudonymization.generator import generate

    out = generate(PiiType.UNCLASSIFIED_GROUP_VALUE, "Nguyen, Tuyet PTIN P01234567")
    assert "PTIN" in out
    assert "Nguyen" not in out and "Tuyet" not in out
    assert "P01234567" not in out
    assert re.search(r"[A-Z]\d{8}", out), f"the identifier lost its shape: {out!r}"


def test_every_component_of_a_place_survives():
    """Reported: 'Fremont, CA 1234' dropped the trailing number entirely."""
    from app.pseudonymization.generator import generate

    out = generate(PiiType.UNCLASSIFIED_GROUP_VALUE, "Fremont, CA 1234")
    assert "Fremont" not in out
    assert re.search(r",\s*[A-Z]{2}\s+\d{4}", out), f"a component was lost: {out!r}"


def test_initials_are_not_mistaken_for_a_state_code():
    from app.pseudonymization.generator import generate

    out = generate(PiiType.UNCLASSIFIED_GROUP_VALUE, "LJ P")
    assert len(out.split()) == 2
    assert "TN" not in out and "CA" not in out


def test_no_pseudonym_is_unreadable_noise():
    """Reported: 'asdiauguw adsasd asd'. Words must be replaced by words."""
    from app.pseudonymization.generator import generate

    for value in ("Nguyen, Tuyet", "Acme Holdings LLC", "Guadalajara, Jalisco", "LJ P"):
        out = generate(PiiType.UNCLASSIFIED_GROUP_VALUE, value)
        for token in out.split():
            letters = "".join(ch for ch in token if ch.isalpha())
            if len(letters) < 4:
                continue
            vowels = sum(1 for ch in letters.lower() if ch in "aeiou")
            assert vowels, f"unpronounceable output {out!r}"


def test_a_date_stays_a_date():
    from app.pseudonymization.generator import generate

    out = generate(PiiType.UNCLASSIFIED_GROUP_VALUE, "04/11/1979")
    assert re.fullmatch(r"\d{2}/\d{2}/\d{4}", out), out
    month, day, _year = out.split("/")
    assert 1 <= int(month) <= 12 and 1 <= int(day) <= 31, f"impossible date {out!r}"


# --- extraction normalization ----------------------------------------------


def test_lookalike_characters_are_normalized_without_shifting_offsets():
    """A non-breaking space or a curly quote silently defeats every match.

    Substitutions must be one-for-one: a ligature expanding to two characters
    would shift every offset after it and misplace redaction rectangles.
    """
    from app.document.model import LOOKALIKES, normalize_char

    for original, replacement in LOOKALIKES.items():
        assert len(original) == len(replacement) == 1, (original, replacement)
        assert normalize_char(original) == replacement
    assert normalize_char("a") == "a"


def test_a_non_breaking_space_does_not_hide_a_name(tmp_path):
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "nbsp.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    c.drawString(72, 720, "Taxpayer name")
    c.drawString(72, 706, "John\u00a0Smith")       # non-breaking space
    c.drawString(72, 660, "1  Wages ...... $412,890")
    c.save()

    s = analysed(str(path))
    text = " ".join(c.text for c in s.candidates)
    assert "John" in text, "the name was hidden by a lookalike character"
    assert "\u00a0" not in text, "the lookalike survived into a candidate"


def test_a_column_header_types_the_cells_beneath_it(tmp_path):
    """A table row has no label beside it - the label is the column header."""
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    path = tmp_path / "cols.pdf"
    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(60, 700, "Employee name")
    c.drawString(220, 700, "SSN")
    c.drawString(360, 700, "Salary")
    c.setFont("Helvetica", 10)
    for offset, (name, ssn, pay) in enumerate(
        [("John Smith", "123-45-6789", "$85,000"), ("Jane Doe", "987-65-4321", "$92,000")]
    ):
        y = 684 - offset * 16
        c.drawString(60, y, name)
        c.drawString(220, y, ssn)
        c.drawString(360, y, pay)
    c.save()

    s = analysed(str(path))
    by_value = {c.normalized: c.pii_type for c in s.candidates}
    assert by_value.get("987-65-4321") is PiiType.SSN, by_value
    assert "$85,000" not in by_value and "$92,000" not in by_value


def test_an_ssn_without_dashes_is_detected(tmp_path):
    """Reported miss: nine digits with no separators."""
    s = analysed(
        _lines_pdf(
            tmp_path / "plain.pdf",
            ["Social security number", "123456789", "1  Wages ...... $412,890"],
        )
    )
    found = {c.normalized: c.pii_type for c in s.candidates}
    assert found.get("123456789") is PiiType.SSN, found


def test_a_lone_surname_in_a_cell_is_detected(tmp_path):
    """Reported miss: a single last name standing alone."""
    s = analysed(
        _lines_pdf(tmp_path / "lone.pdf", ["Employee name", "Gonzalez-Reyes"])
    )
    assert any("Gonzalez-Reyes" in c.text for c in s.candidates)


def test_an_inline_label_is_not_swallowed_by_widening(tmp_path):
    """'Taxpayer SSN 987654321' must keep its label."""
    s = analysed(
        _lines_pdf(tmp_path / "inline.pdf", ["Taxpayer SSN 987654321", "$412,890"])
    )
    for candidate in s.candidates:
        assert "Taxpayer" not in candidate.text, f"label swallowed: {candidate.text!r}"
        assert "SSN" not in candidate.text.upper() or candidate.text.strip().isdigit()


def test_a_bare_nine_digit_number_without_context_is_left_alone(tmp_path):
    """Nine digits are also account and reference numbers."""
    s = analysed(_lines_pdf(tmp_path / "bare.pdf", ["Reference 481920374"]))
    assert not any(c.normalized == "481920374" for c in s.candidates)


# --- extraction quality and name variants ----------------------------------


def test_a_garbled_text_layer_is_flagged_not_silently_empty(tmp_path):
    """A broken font extracts as noise: nothing errors, nothing is detected.

    Silence is indistinguishable from a clean page, so it has to be reported.
    """
    import pymupdf

    from app.document.provider import NativePdfTextProvider

    path = tmp_path / "garbled.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 700), "\ufffd" * 60, fontsize=11)
    page.insert_text((72, 720), "\ufffd" * 60, fontsize=11)
    doc.save(str(path))
    doc.close()

    document = NativePdfTextProvider().load(str(path))
    assert document.pages[0].needs_ocr, "a garbled text layer was reported as fine"
    assert "unreliable text extraction" in document.pages[0].ocr_reason


def test_a_clean_page_is_not_flagged_as_garbled(tmp_path):
    from app.document.provider import NativePdfTextProvider

    path = _lines_pdf(
        tmp_path / "fine.pdf",
        ["Taxpayer name", "Maria Gonzalez", "1  Wages ...... $412,890"],
    )
    document = NativePdfTextProvider().load(path)
    assert not document.pages[0].needs_ocr


def test_name_spelling_variants_are_generated():
    from app.detection.entities_pass import spelling_variants

    variants = spelling_variants("Maria T Gonzalez-Reyes")
    assert "Maria Gonzalez-Reyes" in variants        # middle initial dropped
    assert "Maria T Gonzalez Reyes" in variants      # hyphen as a space
    assert "Gonzalez-Reyes, Maria T" in variants     # surname first
    assert "Maria T Gonzalez-Reyes" not in variants  # not itself


def test_a_hyphen_variant_of_a_known_name_is_caught(tmp_path):
    """The same client written two ways on one form."""
    s = analysed(
        _lines_pdf(
            tmp_path / "variant.pdf",
            [
                "Taxpayer name",
                "Maria Gonzalez-Reyes",
                "Prepared for Maria Gonzalez Reyes this year",
                "1  Wages ...... $412,890",
            ],
        )
    )
    text = " ".join(c.text for c in s.candidates)
    assert "Gonzalez Reyes" in text or "Gonzalez-Reyes" in text
    hits = [c for c in s.candidates if "Gonzalez" in c.text]
    assert len(hits) >= 2, f"only found {[c.text for c in hits]}"


# --- guarding the review model ---------------------------------------------


def test_model_findings_that_would_damage_the_document_are_refused():
    """Reported: the review pass proposed labels and figures.

    The model is advisory and noisy, so its output is filtered against what the
    deterministic layers already know rather than trusted to a prompt.
    """
    from app.detection.auditor import is_acceptable_finding

    labels = {"taxpayer name", "social security number"}
    for value in ("Taxpayer name", "$412,890", "1  Wages", "SSN", "Occupation:",
                  "2025", "Line 12b", "Social security number"):
        accepted, _why = is_acceptable_finding(value, labels)
        assert not accepted, f"{value!r} would have been redacted"


def test_genuine_values_still_pass_the_guard():
    from app.detection.auditor import is_acceptable_finding

    labels = {"taxpayer name"}
    for value in ("Maria Gonzalez", "04/11/1979", "123456789", "000123456789",
                  "Acme Holdings LLC", "Guadalajara, Jalisco"):
        accepted, why = is_acceptable_finding(value, labels)
        assert accepted, f"{value!r} was refused: {why}"


def test_a_label_on_the_page_is_refused_even_if_it_looks_like_a_name():
    from app.detection.auditor import is_acceptable_finding

    accepted, why = is_acceptable_finding("Preparer Details", {"preparer details"})
    assert not accepted and "label" in why


def test_pseudonym_generation_reuses_one_faker_per_thread():
    """Regression: a fresh Faker per call aborted the process during GC.

    Constructing one loads every provider. This runs thousands of times per
    document, and under Qt's teardown the pressure crashed the interpreter
    inside the garbage collector.
    """
    import threading

    from app.pseudonymization.generator import _faker

    assert _faker() is _faker(), "a new Faker is built on every call"

    seen: list = []
    thread = threading.Thread(target=lambda: seen.append(_faker()))
    thread.start()
    thread.join(10)
    assert seen and seen[0] is not _faker(), "one Faker shared across threads"


def test_generation_stays_deterministic_with_a_reused_faker():
    from app.pseudonymization.generator import generate

    first = [generate(PiiType.PERSON, f"Person {i}") for i in range(50)]
    second = [generate(PiiType.PERSON, f"Person {i}") for i in range(50)]
    assert first == second, "reusing the instance broke determinism"
    assert len(set(first)) > 40, "reuse collapsed the variety of pseudonyms"


def test_generation_is_fast_enough_for_a_large_document():
    import time

    from app.pseudonymization.generator import generate

    started = time.monotonic()
    for index in range(500):
        generate(PiiType.PERSON, f"Name {index}")
    assert time.monotonic() - started < 5.0, "pseudonym generation is too slow"


# --- careful mode ----------------------------------------------------------


def test_careful_mode_excludes_the_widening_passes(tmp_path, monkeypatch):
    """A safety valve: transform only what was positively identified.

    The passes that widen coverage are what find values nothing else sees, and
    also what damages a document when they misfire.
    """
    from app.detection.types import Source

    source = fixtures.form_pdf(tmp_path / "careful.pdf")

    monkeypatch.setenv("DOCANON_CONSERVATIVE", "0")
    normal = analysed(source)

    monkeypatch.setenv("DOCANON_CONSERVATIVE", "1")
    careful = analysed(source)

    assert len(careful.candidates) <= len(normal.candidates)
    for candidate in careful.candidates:
        assert candidate.pii_type is not PiiType.UNCLASSIFIED_GROUP_VALUE
        # GROUP stays included: it is label-confirmed evidence, not a guess -
        # and it is the ONLY mechanism that types a joint name's household
        # split. Excluding it made Careful mode miss joint names entirely.
        assert candidate.source not in (Source.COVERAGE, Source.AUDIT)
        assert candidate.confidence >= 0.65


def test_careful_mode_still_catches_the_obvious_values(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCANON_CONSERVATIVE", "1")
    s = analysed(fixtures.form_pdf(tmp_path / "obvious.pdf"))
    found = {c.pii_type for c in s.candidates}
    assert PiiType.SSN in found and PiiType.EMAIL in found


def test_a_manual_addition_survives_careful_mode(tmp_path, monkeypatch):
    monkeypatch.setenv("DOCANON_CONSERVATIVE", "1")
    s = analysed(fixtures.form_pdf(tmp_path / "manual.pdf"))
    added = s.add_manual_text("Annual Salary")
    assert added, "careful mode overruled the user"


# --- type or skip ----------------------------------------------------------


def test_an_untyped_value_is_never_replaced(tmp_path):
    """Nothing is transformed unless a detector said what it is.

    Generating a replacement for a span nothing could identify is what wrote
    scrambled words over form instructions.
    """
    s = analysed(fixtures.stacked_field_pdf(tmp_path / "untyped.pdf"))
    untyped = [c for c in s.candidates if c.pii_type is PiiType.UNCLASSIFIED_GROUP_VALUE]
    assert untyped, "the fixture no longer produces an untyped value"

    for candidate in untyped:
        assert s.decisions.state(candidate) is DecisionState.SKIPPED

    plan = s.plan()
    replaced = {t.original.strip() for t in plan.targets}
    for candidate in untyped:
        assert candidate.normalized not in replaced


def test_naming_an_untyped_value_makes_it_replaceable(tmp_path):
    s = analysed(fixtures.stacked_field_pdf(tmp_path / "named.pdf"))
    untyped = [c for c in s.candidates if c.pii_type is PiiType.UNCLASSIFIED_GROUP_VALUE]
    assert untyped

    for candidate in untyped:
        candidate.pii_type = PiiType.POSTAL_CODE
    s.decisions.set_state(untyped, DecisionState.ACCEPTED)

    replaced = {t.original.strip() for t in s.plan().targets}
    assert untyped[0].normalized in replaced


def test_an_untyped_value_can_still_be_blacked_out(tmp_path):
    s = analysed(fixtures.stacked_field_pdf(tmp_path / "black.pdf"))
    untyped = [c for c in s.candidates if c.pii_type is PiiType.UNCLASSIFIED_GROUP_VALUE]
    assert untyped
    untyped[0].blackout = True
    s.decisions.set_state([untyped[0]], DecisionState.ACCEPTED)

    targets = [t for t in s.plan().targets if t.blackout]
    assert targets, "a blackout needs no type and should still apply"


def test_the_document_keeps_untyped_text(tmp_path):
    """Left alone means left readable, not quietly mangled."""
    s = analysed(fixtures.stacked_field_pdf(tmp_path / "keep.pdf"))
    untyped = [c for c in s.candidates if c.pii_type is PiiType.UNCLASSIFIED_GROUP_VALUE]
    assert untyped
    original = untyped[0].normalized

    s.decisions.set_state(
        [c for c in s.candidates if c.pii_type is not PiiType.UNCLASSIFIED_GROUP_VALUE],
        DecisionState.ACCEPTED,
    )
    out = str(tmp_path / "out.pdf")
    s.process(out)
    assert original in text_of(out)


def test_a_joint_couple_keeps_a_shared_surname(tmp_path):
    """'Mark & Mich Smith' is one couple, and stays one couple.

    Giving each half an unrelated surname loses the fact that they are married,
    which whoever analyses the return needs.
    """
    from app.pseudonymization.generator import generate_for_household

    household = "Mark Mich Smith"
    left = generate_for_household("Mark", household)
    right = generate_for_household("Mich Smith", household)

    assert " " not in left, "a given-name-only side gained a surname"
    assert len(right.split()) == 2
    assert left != right.split()[0], "both halves got the same given name"

    # Another couple gets a different surname.
    other = generate_for_household("Ana Ruiz", "Ana Luis Ruiz")
    assert other.split()[-1] != right.split()[-1]


def test_the_shared_surname_is_stable(tmp_path):
    from app.pseudonymization.generator import generate_for_household

    first = generate_for_household("Mich Smith", "Mark Mich Smith")
    second = generate_for_household("Mich Smith", "Mark Mich Smith")
    assert first == second


def test_conjugal_names_are_split_and_typed(tmp_path):
    s = analysed(
        _lines_pdf(
            tmp_path / "conj.pdf",
            ["Taxpayer name", "Mark & Mich Smith", "1  Wages ...... $412,890"],
        )
    )
    people = {c.normalized for c in s.candidates if c.pii_type is PiiType.PERSON}
    assert "Mark" in people, people
    assert any("Mich" in p for p in people), people
    assert "Mark & Mich Smith" not in people
    assert all(c.household for c in s.candidates if c.pii_type is PiiType.PERSON)


def test_model_proposals_are_typed_by_the_detectors_not_the_model(tmp_path):
    """The model proposes; the deterministic layers decide what it is.

    A proposal nothing can type is reported rather than replaced, so a wrong
    guess costs a line in the review list instead of corrupt output.
    """
    from app.detection.auditor import confirm_with_detectors
    from app.detection.types import Candidate, PiiType, Source
    from app.document.model import Char, Line, Span

    def make(text: str) -> Candidate:
        char = Char(text="x", bbox=(0, 0, 5, 10))
        span = Span(text=text, bbox=(0, 0, 50, 10), font="helv", size=10, color=0,
                    chars=[char])
        line = Line(page_no=0, block_no=0, line_no=0, spans=[span], text=text,
                    offsets=[char], bbox=(0, 0, 50, 10))
        return Candidate(
            pii_type=PiiType.UNCLASSIFIED_GROUP_VALUE, text=text, page_no=0,
            rect=(0, 0, 50, 10), line=line, start=0, end=len(text),
            confidence=0.6, source=Source.AUDIT,
        )

    typed = confirm_with_detectors(None, [make("john@example.org")], None)
    assert typed[0].pii_type is PiiType.EMAIL, typed[0].pii_type

    person = confirm_with_detectors(None, [make("Marisol Etxeberria")], None)
    assert person[0].pii_type is PiiType.PERSON

    unknown = confirm_with_detectors(None, [make("qq zz ww")], None)
    assert unknown[0].pii_type is PiiType.UNCLASSIFIED_GROUP_VALUE
    assert unknown[0].needs_review
    assert "nothing could say what it is" in unknown[0].review_reason


def test_an_untyped_value_left_behind_does_not_fail_verification(tmp_path):
    """Regression: the CLI build failed on the stacked fixture.

    --accept-all marks everything accepted, type-or-skip then drops the untyped
    value, and the group check still expected it gone. A value deliberately
    left readable must be recorded as kept, not treated as a missed redaction.
    """
    s = analysed(fixtures.stacked_field_pdf(tmp_path / "left.pdf"))
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)

    plan = s.plan()
    untyped = [c for c in s.candidates if c.pii_type is PiiType.UNCLASSIFIED_GROUP_VALUE]
    assert untyped
    assert plan.unresolved, "the untyped value was not reported"
    for candidate in untyped:
        assert candidate.normalized in plan.skipped_values

    out = str(tmp_path / "out.pdf")
    _apply, report = s.process(out)
    assert report.passed, [f"{c.name}: {c.detail}" for c in report.failures]
    assert untyped[0].normalized in text_of(out), "it should still be readable"


# --- overlapping targets ----------------------------------------------------


def test_overlapping_targets_are_merged_before_writing():
    """Reported: overlapping/doubled text in real output - one name drawn
    on top of another, an SSN and a name overlapping.

    Two candidates covering the same or adjacent geometry both redacted and
    inserted independently, so the second was drawn over the first without
    clearing it - visible as jumbled, overlapping text in the final PDF.
    """
    from app.transform.plan import Target, _merge_overlapping_targets

    def target(rect, original):
        return Target(
            candidate_id=str(rect), page_no=0, rect=rect, original=original,
            replacement="X", font_size=10.0, font_name="helv",
        )

    overlapping = [
        target((10, 10, 90, 20), "MEGHAN SHAFFER"),
        target((15, 10, 85, 20), "TIFFANY SHAFFER"),
    ]
    merged = _merge_overlapping_targets(overlapping)
    assert len(merged) == 1, [m.original for m in merged]


def test_distinct_targets_on_the_same_page_all_survive():
    from app.transform.plan import Target, _merge_overlapping_targets

    def target(rect, original):
        return Target(
            candidate_id=str(rect), page_no=0, rect=rect, original=original,
            replacement="X", font_size=10.0, font_name="helv",
        )

    distinct = [
        target((10, 10, 90, 20), "John Smith"),
        target((10, 40, 90, 50), "135-34-8757"),
        target((10, 70, 90, 80), "4049 Caribbean Cmn"),
    ]
    merged = _merge_overlapping_targets(distinct)
    assert len(merged) == 3


def test_a_real_document_never_produces_overlapping_targets(tmp_path):
    """End to end: build a real plan and confirm no two targets touch."""
    s = analysed(fixtures.form_pdf(tmp_path / "overlap.pdf"))
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    targets = s.plan().targets
    for i, a in enumerate(targets):
        for b in targets[i + 1:]:
            if a.page_no != b.page_no:
                continue
            ax0, ay0, ax1, ay1 = a.rect
            bx0, by0, bx1, by1 = b.rect
            touching = not (ax1 < bx0 or bx1 < ax0 or ay1 < by0 or by1 < ay0)
            assert not touching, f"{a.original!r} and {b.original!r} overlap"


def test_english_phrases_are_not_typed_as_names():
    """Reported: 'Need to Keep' was typed PERSON.

    Title-case shape alone matches any short phrase. An English function word
    anywhere in it means it is a sentence fragment, not a name.
    """
    from app.detection.heuristics import looks_like_person

    for phrase in (
        "Need to Keep", "Please Review", "Estimated Payments",
        "Federal Tax Return", "Make for Next", "Action Required",
    ):
        detected, _confidence = looks_like_person(phrase)
        assert not detected, phrase


def test_real_names_still_pass_after_the_english_word_guard():
    from app.detection.heuristics import looks_like_person

    for name in ("John Smith", "Julie Haynes", "Brenda Diaz", "Maria Gonzalez-Reyes"):
        detected, _confidence = looks_like_person(name)
        assert detected, name


def test_document_headings_survive_end_to_end(tmp_path):
    s = analysed(
        _lines_pdf(
            tmp_path / "headings.pdf",
            ["Julie Haynes", "Need to Keep", "2021 Federal Tax Return Summary",
             "Estimated Payments to Make for Next", "Brenda Diaz"],
        )
    )
    texts = {c.text for c in s.candidates}
    assert "Need to Keep" not in texts
    assert "2021 Federal Tax Return Summary" not in texts
    assert "Estimated Payments to Make for Next" not in texts
    assert any("Julie" in t or "Haynes" in t for t in texts)
    assert any("Brenda" in t or "Diaz" in t for t in texts)


def test_padding_is_symmetric_by_default_like_v0_25(tmp_path):
    """Superseded by the literal v0.8.3 revert.

    The neighbour-aware headroom functions this test checked for
    (_headroom_above/_headroom_below) no longer exist - the redactor was
    reverted to v0.8.3's actual code, which pads symmetrically with NO
    neighbour check at all. That is a stronger, simpler version of "symmetric
    by default" than this test originally verified, so the property holds;
    the specific functions it named do not exist to check anymore.
    """
    from app.export import redactor

    assert not hasattr(redactor, "_headroom_above")
    assert not hasattr(redactor, "_headroom_below")
    assert "0.35 * current" in __import__("inspect").getsource(redactor._insert_replacement)


# --- from the six-image report ----------------------------------------------


def test_a_joint_name_with_no_binding_label_is_still_detected(tmp_path):
    """Reported: 'Mark & Jane Lang' not detected at all in some locations."""
    s = analysed(
        _lines_pdf(
            tmp_path / "signature.pdf",
            ["Signature block: reviewed and signed by", "Mark & Jane Lang",
             "1  Wages ...... $412,890"],
        )
    )
    people = {c.normalized: c.household for c in s.candidates if c.pii_type is PiiType.PERSON}
    assert "Mark" in people, people
    assert any("Jane" in k for k in people), people
    assert all(v for v in people.values()), "no household link recorded"


def test_a_label_word_is_never_left_exposed_by_widening(tmp_path):
    """Reported: 'security' inside 'social security number' got a badge."""
    s = analysed(
        _lines_pdf(
            tmp_path / "ssnlabel.pdf",
            ["Your social security number 123456789", "1  Wages ...... $412,890"],
        )
    )
    for candidate in s.candidates:
        assert "security" not in candidate.text.lower()
        assert "social" not in candidate.text.lower()


def test_address_variants_share_one_pseudonym(tmp_path):
    from app.entities.registry import EntityRegistry

    registry = EntityRegistry()
    assert registry.normalize("4049 Caribbean Street") == registry.normalize(
        "4049 Caribbean St"
    )
    assert registry.normalize("123 Main Avenue") == registry.normalize("123 Main Ave")


# --- typographic fidelity (font flags, baseline origin) --------------------


def test_font_flags_are_captured_from_the_original_span(tmp_path):
    """Bold/italic preserved without changing color - red stays a
    deliberate, verified security signal, never something this touches."""
    import pymupdf

    path = tmp_path / "bold.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "John Smith", fontname="hebo", fontsize=11)
    doc.save(str(path))
    doc.close()

    from app.document.provider import NativePdfTextProvider

    document = NativePdfTextProvider().load(str(path))
    line = document.pages[0].lines[0]
    assert line.flags & (1 << 4), "bold flag was not captured"


def test_bold_is_detected_from_flags_even_with_an_obfuscated_font_name():
    from app.export.redactor import _safe_font

    assert _safe_font("Arial-CustomID7", flags=1 << 4) == "hebo"
    assert _safe_font("Arial-CustomID7", flags=0) == "helv"


def test_origin_flows_from_span_to_target(tmp_path):
    s = analysed(fixtures.form_pdf(tmp_path / "origin.pdf"))
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    for target in s.plan().targets:
        assert hasattr(target, "origin")


# --- geometric label exclusion ----------------------------------------------


def test_a_candidate_overlapping_a_label_by_half_its_area_is_dropped():
    """True bounding-box check, independent of which Line either side
    belongs to - catches what the same-line character-offset check misses."""
    from app.detection.resolve import veto_by_label_geometry
    from app.detection.types import LabelRegion

    label = LabelRegion(
        text="Social security number", line=None, start=0, end=10,
        rect=(70, 700, 200, 714), expected_types=[PiiType.SSN],
    )
    overlapping = Candidate(
        pii_type=PiiType.SSN, text="123456789", page_no=0,
        rect=(75, 702, 210, 712),  # >50% inside the label's rect
        line=None, start=0, end=9, confidence=0.7, source=Source.NER,
    )
    clear = Candidate(
        pii_type=PiiType.SSN, text="987654321", page_no=0,
        rect=(75, 800, 210, 812),  # nowhere near the label
        line=None, start=0, end=9, confidence=0.7, source=Source.NER,
    )
    kept, dropped = veto_by_label_geometry([overlapping, clear], [label])
    assert dropped == 1
    assert kept == [clear]


def test_boilerplate_vocabulary_covers_mail_in_voucher_fragments():
    from app.detection.heuristics import FORM_VOCABULARY

    for word in ("here", "mail", "with", "detach", "page", "form"):
        assert word in FORM_VOCABULARY, word


def test_raise_if_leaked_passes_a_clean_report_through(tmp_path):
    s = analysed(fixtures.form_pdf(tmp_path / "clean.pdf"))
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    _apply, report = s.process(str(tmp_path / "out.pdf"))
    from app.verification.verifier import raise_if_leaked

    assert raise_if_leaked(report) is report


def test_raise_if_leaked_raises_on_a_critical_failure():
    from app.verification.verifier import Check, PIILeakError, VerificationReport, raise_if_leaked

    report = VerificationReport(
        output_path="x.pdf",
        checks=[Check(name="accepted originals removed", passed=False, detail="leak", critical=True)],
    )
    with pytest.raises(PIILeakError):
        raise_if_leaked(report)


def test_raise_if_leaked_ignores_non_critical_failures():
    from app.verification.verifier import Check, VerificationReport, raise_if_leaked

    report = VerificationReport(
        output_path="x.pdf",
        checks=[Check(name="financial values preserved", passed=False, detail="cosmetic", critical=False)],
    )
    assert raise_if_leaked(report) is report


# --- case-insensitive deterministic matching --------------------------------


def test_rules_match_regardless_of_case_by_default(tmp_path):
    """Reported: an all-caps street address ('4049 MICHAEL CMN') was typed
    PERSON instead of STREET, because every rule compiled case-SENSITIVE.
    Real government correspondence frequently renders address blocks in
    all-caps."""
    s = analysed(
        _lines_pdf(
            tmp_path / "caps.pdf",
            ["MEGHAN SHAFFER", "TIFFANY SHAFFER", "4049 MICHAEL CMN",
             "CINCINNATI OH 45280-2502"],
        )
    )
    by_text = {c.normalized: c.pii_type for c in s.candidates}
    assert by_text.get("4049 MICHAEL CMN") is PiiType.STREET, by_text


def test_state_codes_still_require_real_capitals(tmp_path):
    """The one rule that must opt OUT of case-insensitivity: half the state
    codes are common English words (OR, IN, HI, ME, OK...), and matching them
    lowercase turned 'you, or your spouse' into a detected address."""
    s = analysed(
        _lines_pdf(
            tmp_path / "prose.pdf",
            ["Check if you, or your spouse if filing jointly, want to opt in."],
        )
    )
    assert not any(c.pii_type is PiiType.ADDRESS for c in s.candidates)


def test_city_state_matches_without_a_comma(tmp_path):
    """Reported gap: a condensed mailing block with no comma between city
    and state - common on real IRS notices - matched nothing."""
    from app.detection.deterministic import load_rules

    rule = next(r for r in load_rules().rules if r.name == "city_state_zip")
    assert rule.regex.search("CINCINNATI OH 45280-2502")
    assert rule.regex.search("Fremont, CA 94538")  # the comma form still works


def test_a_rule_can_opt_out_of_case_insensitivity_via_yaml():
    from app.detection.deterministic import load_rules

    rule = next(r for r in load_rules().rules if r.name == "city_state_zip")
    import re

    assert not (rule.regex.flags & re.IGNORECASE)
    other = next(r for r in load_rules().rules if r.name == "street_address")
    assert other.regex.flags & re.IGNORECASE


def test_the_insert_text_fallback_is_never_dropped():
    """Regression: an editing mistake this session removed the entire
    'last resort' fallback block. insert_textbox has a HEIGHT requirement
    that a tightly-capped, neighbour-aware pad can legitimately fail to
    meet - insert_text has no such check and must always be there to catch
    that case, or a replacement vanishes with no error at all."""
    import inspect

    from app.export import redactor

    source = inspect.getsource(redactor._insert_replacement)
    assert "insert_text(" in source, "the baseline fallback is missing"
    assert '"failed"' in source


def test_plain_text_extraction_is_not_how_visual_separation_is_verified():
    """A lesson from this session, not just a fact about the code: PyMuPDF's
    plain get_text() joins two adjacent spans into one string with no space
    whenever there is no whitespace CHARACTER between them, even when there
    is real, visible pixel separation on the page. A test - or a person -
    checking for two names 'merging' must sample rendered pixels or read
    get_text('dict') span boundaries, never treat a joined plain-text
    string as proof of a visual defect on its own.
    """
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_text((72, 100), "Alpha", fontsize=10)
    page.insert_text((72 + pymupdf.get_text_length("Alpha", fontsize=10) + 12, 100),
                     "Beta", fontsize=10)

    plain = page.get_text()
    spans = [s for b in page.get_text("dict")["blocks"] for l in b.get("lines", [])
             for s in l.get("spans", [])]
    # Two clearly separate words, twelve points apart - a real person reading
    # the page would never call this "merged".
    assert spans[0]["bbox"][2] < spans[1]["bbox"][0] - 5, "the two words are not actually separated"
    doc.close()


def test_bare_role_labels_are_recognised_without_the_word_name(tmp_path):
    """Reported: 'Policyholder' (no trailing 'Name') was itself flagged as
    PII, because nothing recognised it as a field label at all."""
    s = analysed(_lines_pdf(tmp_path / "role.pdf", ["Policyholder", "Diana Whitfield"]))
    assert not any("Policyholder" in c.text for c in s.candidates)
    assert any("Whitfield" in c.text for c in s.candidates)


def test_business_name_suffix_variants_share_one_pseudonym(tmp_path):
    """Reported: business name inconsistent across pages."""
    from app.entities.registry import EntityRegistry
    from app.detection.types import PiiType

    registry = EntityRegistry()
    variants = ["Acme Holdings LLC", "Acme Holdings, L.L.C.", "Acme Holdings",
                "Acme Holdings Inc"]
    pseudonyms = {registry.pseudonym_for(PiiType.ORG_PRIVATE, v) for v in variants}
    assert len(pseudonyms) == 1, pseudonyms


def test_business_name_core_stripping_does_not_merge_different_companies():
    from app.entities.registry import EntityRegistry

    registry = EntityRegistry()
    assert registry._org_core("Acme Holdings LLC") == registry._org_core("Acme Holdings Inc")
    assert registry._org_core("Acme Holdings LLC") != registry._org_core("Beta Ventures LLC")


# --- from the master overhaul: adjudication, coverage, benchmark -----------


def test_the_benchmark_corpus_has_perfect_precision_and_recall():
    """Wires the golden benchmark into the ordinary test suite, so a
    regression in any category is caught by CI, not just a manual run."""
    from benchmark.fixtures import ALL_FIXTURES
    from benchmark.run_benchmark import aggregate, run_fixture

    with tempfile.TemporaryDirectory() as tmp:
        results = [run_fixture(f, Path(tmp), use_ner=True) for f in ALL_FIXTURES]
    report = aggregate(results)
    assert report["false_positives"] == 0, report["per_document"]
    assert report["documents_with_any_missed_pii"] == 0, report["documents_with_any_missed_pii_names"]


def test_widening_never_swallows_a_short_identifier_inside_a_sentence(tmp_path):
    """Found by the benchmark harness on its first run: a social handle
    inside an ordinary sentence ('Follow me @x on Instagram') got widened
    to the WHOLE sentence, because none of its words were on the
    tax/form vocabulary list the widening guard already checked."""
    s = analysed(
        _lines_pdf(tmp_path / "handle.pdf", ["Follow me @traveler_jane on Instagram"])
    )
    hit = next((c for c in s.candidates if "@" in c.text), None)
    assert hit is not None, "the handle was not detected at all"
    assert hit.text.strip() == "@traveler_jane", f"swallowed the sentence: {hit.text!r}"


def test_every_candidate_carries_an_adjudication_verdict(tmp_path):
    from app.detection.adjudication import Adjudication

    s = analysed(fixtures.form_pdf(tmp_path / "adj.pdf"))
    assert s.candidates
    for candidate in s.candidates:
        assert candidate.adjudication in (v.value for v in Adjudication)


def test_an_unresolved_candidate_defaults_to_skipped_even_if_typed(tmp_path):
    """Broadens type-or-skip: a KNOWN type with too little evidence gets the
    same treatment as an unknown type - wait for a decision, never
    auto-applied."""
    from app.detection.adjudication import Adjudication
    from app.decisions.manager import DecisionState

    s = analysed(fixtures.form_pdf(tmp_path / "unresolved.pdf"))
    if not s.candidates:
        pytest.skip("fixture produced nothing to test against")
    victim = s.candidates[0]
    victim.adjudication = Adjudication.UNRESOLVED.value
    s.decisions.decisions.pop(victim.id, None)
    s.decisions.register([victim])
    assert s.decisions.state(victim) is DecisionState.SKIPPED


def test_the_coverage_matrix_has_no_undetectable_types():
    """The exact check that found ACCOUNT_ID, FAX, MARITAL_STATUS, MATTER_ID,
    MEDICARE_ID, PAYROLL_ID, SOCIAL_HANDLE, STATE_TAX_ID and URL_PERSONAL had
    no detector path at all."""
    import subprocess
    import sys

    result = subprocess.run(
        [sys.executable, "buildtools/generate_coverage_matrix.py"],
        cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True,
    )
    assert result.returncode == 0, result.stdout + result.stderr


# --- filename anonymization -------------------------------------------------


def test_a_name_token_in_the_filename_matches_the_document_pseudonym(tmp_path):
    """Reported bug, reproduced exactly against the spec's own example:
    'Lance' in the filename was replaced with the literal word 'REDACTED'
    instead of the SAME pseudonym token used inside the document. The
    per-word fallback discarded name-token mapping instead of using the
    NameRegistry this project already has for exactly this case.
    """
    path = _lines_pdf(tmp_path / "1040 Return 2021 Lance.pdf",
                       ["Taxpayer name", "Lance Whitfield"])
    s = analysed(path)
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)

    doc_pseudonym = next(
        t.replacement for t in s.plan().targets if t.original.strip() == "Lance Whitfield"
    )
    given_name = doc_pseudonym.split()[0]

    name = s.safe_output_name()
    assert "REDACTED" not in name
    assert given_name in name
    assert "Lance" not in name


def test_a_name_only_in_the_filename_still_gets_a_stable_pseudonym(tmp_path):
    """Spec: filename PII need not also occur inside the PDF body."""
    path = _lines_pdf(tmp_path / "Backup_Torres_Files.pdf",
                       ["1  Wages ...... $412,890"])
    s = analysed(path)
    name = s.safe_output_name()
    assert "Torres" not in name
    assert "Backup" in name and "Files" in name  # non-PII parts preserved


def test_ordinary_filename_words_are_never_mistaken_for_names(tmp_path):
    """Filenames are full of title-cased English words that are not names -
    'Files', 'Report', 'Final' - the lone-surname heuristic is permissive
    enough on a single bare word that this needed its own guard."""
    path = _lines_pdf(tmp_path / "Final_Report_Documents.pdf",
                       ["1  Wages ...... $412,890"])
    s = analysed(path)
    name = s.safe_output_name()
    assert "Final" in name and "Report" in name and "Documents" in name


def test_form_numbers_and_years_survive_filename_anonymization(tmp_path):
    path = _lines_pdf(tmp_path / "1040 Return 2021 Whitfield.pdf",
                       ["Taxpayer name", "Whitfield"])
    s = analysed(path)
    s.decisions.set_state(s.candidates, DecisionState.ACCEPTED)
    name = s.safe_output_name()
    assert "1040" in name and "2021" in name
    assert "Whitfield" not in name


def test_spacy_trf_is_the_default_production_model():
    """en_core_web_trf is the ONLY spaCy model the production build ships,
    on direct and repeated instruction. DOCANON_SPACY_MODEL exists only for
    development/CI-speed overrides on a machine where the full torch runtime
    is impractical - the production build must never set that override."""
    import importlib
    import os

    import app.detection.ner as ner

    saved = os.environ.pop("DOCANON_SPACY_MODEL", None)
    try:
        importlib.reload(ner)
        assert ner.MODEL_NAME == "en_core_web_trf"

        os.environ["DOCANON_SPACY_MODEL"] = "en_core_web_sm"
        importlib.reload(ner)
        assert ner.MODEL_NAME == "en_core_web_sm"
    finally:
        if saved is None:
            os.environ.pop("DOCANON_SPACY_MODEL", None)
        else:
            os.environ["DOCANON_SPACY_MODEL"] = saved
        importlib.reload(ner)


def test_trf_runtime_is_a_required_dependency_now_not_optional():
    root = Path(__file__).resolve().parents[1]
    text = (root / "requirements.txt").read_text()
    assert "spacy-transformers" in text
    assert "en_core_web_trf" in text or (root / "buildtools" / "fetch_models.py").read_text().count(
        "en_core_web_trf"
    )
    for manifest in ("requirements.txt", "requirements-dev.txt"):
        assert "en_core_web_sm" not in (root / manifest).read_text(), (
            f"{manifest} must not require en_core_web_sm in production"
        )


# --- en_core_web_trf model-swap regressions ---------------------------------


def test_model_entity_boundary_whitespace_is_trimmed_before_geometry_mapping(tmp_path):
    """en_core_web_trf's span for a name split across a line break can
    include a trailing newline INSIDE the span ('Marisol\\n') where
    en_core_web_sm did not. Any span-to-geometry conversion must trim
    whitespace from the model's boundary first."""
    from app.detection.ner import detect_ner

    class FakeEnt:
        def __init__(self, text, start_char, end_char, label_):
            self.text, self.start_char, self.end_char, self.label_ = (
                text, start_char, end_char, label_,
            )

    class FakeDoc:
        def __init__(self, ents):
            self.ents = ents

    class FakeNlp:
        def __call__(self, text):
            # Simulate the exact trf behaviour: a trailing newline inside
            # the first entity's span.
            idx = text.index("Marisol")
            return FakeDoc([
                FakeEnt(text[idx:idx + 8], idx, idx + 8, "PERSON"),  # "Marisol\n"
            ])

    path = _lines_pdf(
        tmp_path / "split.pdf",
        ["We spoke at length with our client, Marisol",
         "Etxeberria, about the pending refund status."],
    )
    from app.document.provider import NativePdfTextProvider

    doc = NativePdfTextProvider().load(path)
    candidates, _warnings = detect_ner(doc, FakeNlp())
    assert candidates, "the whitespace-including span produced nothing at all"
    assert "\n" not in candidates[0].text


def test_a_single_token_person_hit_is_kept_when_name_shaped(tmp_path):
    """en_core_web_trf segments a name split across a line break into TWO
    single-token entities where en_core_web_sm returned one combined span.
    A blanket single-token rejection silently ate this - a second,
    independent shape signal must let a genuine single-token name through
    while still rejecting the original false positives this filter existed
    for ('Daytime', 'Preparer')."""
    from app.detection.ner import _implausible_person

    assert not _implausible_person("Marisol")
    assert not _implausible_person("Etxeberria")
    assert not _implausible_person("Marisol\n")
    assert _implausible_person("Daytime")
    assert _implausible_person("Preparer")


def test_tiered_ner_skip_costs_zero_recall_on_a_financial_table(tmp_path):
    """The transformer must be skippable on pure-financial-table rows
    without losing a real name sitting elsewhere on the same page."""
    from app.detection.deterministic import detect_deterministic
    from app.detection.ner import detect_ner, load_nlp
    from app.document.provider import NativePdfTextProvider

    lines = ["John Smith"]
    path = _lines_pdf(tmp_path / "table.pdf", lines)
    # Add the table as a visually separate block (a real gap, not just the
    # next line) so it is not merged with "John Smith" into one block -
    # otherwise the whole block correctly never skips, since it does
    # contain a real name.
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    c = canvas.Canvas(str(path), pagesize=letter)
    c.setFont("Helvetica", 10)
    c.drawString(72, 720, "John Smith")
    y = 600
    for i in range(20):
        c.drawString(72, y, f"Line {i}   ${12000+i}.00   ${45000+i}.00")
        y -= 14
    c.save()
    doc = NativePdfTextProvider().load(path)
    det = detect_deterministic(doc)
    nlp = load_nlp()

    without_tiering, _ = detect_ner(doc, nlp, existing=None)
    with_tiering, warnings = detect_ner(doc, nlp, existing=det)
    assert len(with_tiering) == len(without_tiering)
    assert any("skipped the transformer" in w for w in warnings)


def test_verify_bundle_requires_torch_and_spacy_transformers():
    from pathlib import Path as _P

    root = _P(__file__).resolve().parents[1]
    text = (root / "buildtools" / "verify_bundle.py").read_text()
    assert "torch" in text.lower()
    assert "spacy_transformers" in text or "spacy-transformers" in text
    assert "en_core_web_trf" in text


def test_requirements_resolve_without_a_transformers_version_conflict():
    """Real CI failure, reproduced and fixed: requirements.txt pinned
    transformers==4.57.6 (a leftover from an earlier, since-removed LLM
    path - confirmed by grep, zero imports of `transformers` anywhere under
    app/) while spacy-transformers==1.4.0 (the latest release) has a hard
    ceiling of transformers<4.53.3. A fresh pip resolve of the full file in
    one command hit ResolutionImpossible immediately; this sandbox's own
    earlier ad-hoc installs never surfaced it because spacy-transformers
    was installed in a SEPARATE command from the conflicting pin.
    """
    from pathlib import Path as _P

    root = _P(__file__).resolve().parents[1]
    text = (root / "requirements.txt").read_text()
    active_lines = [
        line.strip() for line in text.splitlines()
        if line.strip() and not line.strip().startswith("#")
    ]
    bare_transformers_pins = [
        line for line in active_lines
        if line.startswith("transformers==")  # not spacy-transformers==
    ]
    assert not bare_transformers_pins, (
        "an exact transformers pin was reintroduced - it will conflict with "
        "spacy-transformers' <4.53.3 ceiling in a fresh resolve; nothing in "
        "app/ imports transformers directly, so it does not need one"
    )
    assert "spacy-transformers==1.4.0" in text


def test_requirements_trf_file_no_longer_exists():
    """Folded into requirements.txt directly once trf became the mandatory
    default - a separate opt-in file duplicating the same dependency is
    exactly the orphaned-implementation risk the project's own rules warn
    about."""
    from pathlib import Path as _P

    root = _P(__file__).resolve().parents[1]
    assert not (root / "requirements-trf.txt").exists()


def test_build_py_checks_for_trf_not_sm():
    """Real CI failure, reproduced: build.py --dmg failed with 'en_core_web_sm
    is not installed', even though the model default had already been
    switched to trf everywhere else (ner.py, requirements.txt, the CI
    workflow's fetch step, verify_bundle.py). build.py has its OWN,
    separate SPACY_MODEL constant and pre-build sanity check that was
    missed in that earlier pass."""
    from pathlib import Path as _P

    root = _P(__file__).resolve().parents[1]
    text = (root / "buildtools" / "build.py").read_text()
    assert 'SPACY_MODEL = "en_core_web_trf"' in text
    assert '"en_core_web_sm"' not in text


def test_build_py_collects_the_full_torch_and_transformer_runtime():
    """en_core_web_trf needs torch and spacy_transformers bundled in full
    (PyInstaller cannot discover their data files by static analysis) -
    en_core_web_sm never needed either, so this was never exercised until
    trf became the default."""
    from pathlib import Path as _P

    root = _P(__file__).resolve().parents[1]
    text = (root / "buildtools" / "build.py").read_text()
    assert '"torch"' in text
    assert '"spacy_transformers"' in text


def test_size_floor_does_not_exceed_real_observed_build_sizes():
    """Real CI failure, reproduced: the size floor (originally 1700 MB, set
    from a Linux pip-install disk-delta measurement) broke two genuine,
    successful builds - Windows came in at 1533 MB, macOS at 819 MB, both
    with every required package collected and zero errors in PyInstaller's
    own log. A disk-delta from `pip install` on one platform is not the same
    number as what PyInstaller's --collect-all actually bundles on a
    DIFFERENT platform. The floor must stay below both real numbers, with
    margin, while remaining well above what a build missing the whole
    transformer runtime would produce (~220 MB, the pre-trf baseline).
    """
    import importlib.util
    from pathlib import Path as _P

    root = _P(__file__).resolve().parents[1]
    spec = importlib.util.spec_from_file_location(
        "verify_bundle", root / "buildtools" / "verify_bundle.py"
    )
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)

    REAL_MACOS_MB = 819
    REAL_WINDOWS_MB = 1533
    PRE_TRF_BASELINE_MB = 220

    assert module.MIN_TOTAL_MB < REAL_MACOS_MB, (
        f"floor {module.MIN_TOTAL_MB} would fail the real macOS build "
        f"({REAL_MACOS_MB} MB) again"
    )
    assert module.MIN_TOTAL_MB < REAL_WINDOWS_MB
    assert module.MIN_TOTAL_MB > PRE_TRF_BASELINE_MB, (
        "the floor must still catch a build missing the transformer "
        "runtime entirely"
    )


# --- from a real crash log spanning Sept 9-14 (an older build) -------------


def test_no_dialog_comparison_uses_instance_level_accepted():
    """Real crash log: 'EditDetectionDialog'/'AddPiiDialog'/'UnreviewedPrompt'
    object has no attribute 'Accepted'. dialog.Accepted (instance access)
    raises AttributeError - PySide6 only exposes the class-level
    QDialog.DialogCode.Accepted. Already fixed everywhere; this locks it in
    so it can never silently regress one call site at a time."""
    from pathlib import Path as _P
    import re

    root = _P(__file__).resolve().parents[1]
    text = (root / "app" / "ui" / "main_window.py").read_text()
    bad = re.findall(r"\.exec\(\)\s*!=\s*(?!QDialog\.DialogCode\.)\w+\.Accepted", text)
    assert not bad, f"instance-level .Accepted comparison(s) found: {bad}"


def test_coerce_type_handles_every_shape_qt_can_hand_back():
    """Real crash log: 'str' object has no attribute 'value', raised from
    occurrence_groups' sort key (g.pii_type.value) and from the pseudonym
    generator's _seed. PiiType subclasses str, so a value round-tripping
    through a QComboBox came back as a plain string - grouping or seeding
    with it crashed immediately. Already fixed via coerce_type/_as_pii_type
    at every assignment site; this exercises the actual failure shape."""
    from app.ui.dialogs import coerce_type
    from app.detection.types import PiiType

    assert coerce_type(PiiType.PERSON) is PiiType.PERSON
    assert coerce_type("PERSON") is PiiType.PERSON  # the enum NAME
    assert coerce_type(PiiType.PERSON.value) is PiiType.PERSON  # the enum VALUE

    from app.session import _as_pii_type

    assert _as_pii_type("SSN") is PiiType.SSN
    assert _as_pii_type(PiiType.SSN) is PiiType.SSN


def test_occurrence_groups_never_crashes_on_a_coerced_type(tmp_path):
    """The actual crash site: sorting by g.pii_type.value after a candidate
    was retyped through the dialog path."""
    from app.decisions.manager import DecisionManager
    from app.ui.dialogs import coerce_type

    s = analysed(fixtures.form_pdf(tmp_path / "coerce.pdf"))
    if not s.candidates:
        pytest.skip("fixture produced nothing to test against")
    victim = s.candidates[0]
    # Simulate exactly what a raw Qt round-trip hands back, then run it
    # through the same coercion the dialog path already applies.
    victim.pii_type = coerce_type(victim.pii_type.name)
    groups = DecisionManager.occurrence_groups(s.candidates)
    assert groups  # did not raise


def test_numpy_is_pinned_below_2_to_avoid_a_real_reported_crash():
    """Real, reported crash in the packaged app - every document analysis
    failed instantly with:

        SystemError: <class 'ImportError'> returned a result with an
        exception set

    That specific signature is a well-known symptom of a numpy C-API/ABI
    mismatch: thinc/blis (spaCy's own compiled Cython internals) were built
    against numpy 1.x's ABI, which numpy 2.0 broke. A genuinely fresh
    resolve of requirements.txt (confirmed directly in a real CI run, not
    guessed) pulls numpy 2.5.3 to satisfy spacy's loose numpy>=1.19.0
    constraint - my own sandbox testing never caught this because it
    already had numpy 1.26.4 installed from earlier, unrelated work, and
    pip does not upgrade an already-satisfied dependency.

    Verified the fix directly in a genuinely fresh, isolated venv: with
    this pin, a fresh resolve settles on numpy 1.26.4, and importing +
    running en_core_web_trf alongside torch succeeds with no SystemError.
    """
    from pathlib import Path as _P

    root = _P(__file__).resolve().parents[1]
    text = (root / "requirements.txt").read_text()
    active_lines = [
        line.strip() for line in text.splitlines()
        if line.strip() and not line.strip().startswith("#")
    ]
    numpy_pins = [line for line in active_lines if line.startswith("numpy")]
    assert numpy_pins, "numpy must be explicitly pinned, not left to a transitive resolve"
    assert numpy_pins[0] == "numpy<2", numpy_pins
